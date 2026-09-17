"""Company isolation, original-preserving imports and continuity at the public seams."""
import base64
import copy
import io
import json
from pathlib import Path
import pytest
from docx import Document
from fastapi.testclient import TestClient
from fastapi import HTTPException
from backend.app import create_app
from backend import announcement_history as history, file_ingestion, library_admin, template_authoring
from backend.domain import Seeds
from backend.gates import fingerprint


def docx_bytes(code='300001',title='第三届董事会第二次会议决议公告',day='2025年6月3日',announcement_number='2025-002'):
    doc=Document();doc.add_paragraph(title,style='Title')
    doc.add_paragraph(f'证券代码：{code}  公告编号：{announcement_number}')
    doc.add_paragraph(f'本次会议于{day}召开。')
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='议案';table.cell(0,1).text='结果';table.cell(1,0).text='经营计划';table.cell(1,1).text='通过'
    doc.add_paragraph('文末正文不可丢失。');out=io.BytesIO();doc.save(out);return out.getvalue()


@pytest.fixture
def workspace(tmp_path):
    public=tmp_path/'data/public/boards/chinext';public.mkdir(parents=True)
    for name,value in [('catalog.json',{'board':'chinext','sources':[],'cases':[]}),('profiles.json',[]),('rules.json',[]),('instruments.json',[])]:
        (public/name).write_text(json.dumps(value))
    templates=tmp_path/'templates/boards/chinext';templates.mkdir(parents=True)
    (templates/'manifest.json').write_text('[]');(templates/'layout_profiles.json').write_text('[]')
    app=create_app(tmp_path/'var',tmp_path,{'allowed_hosts':['testserver']})
    with TestClient(app) as c:
        token=c.get('/api/session').json()['csrf_token'];c.headers.update({'Origin':'http://testserver','X-CSRF-Token':token})
        for code in ('300001','300002'):history.register_company(tmp_path,'chinext',code,'测试公司'+code)
        yield c,tmp_path,app


def upload(c,raw=None,company='300001',collection='history'):
    r=c.post('/api/library/imports',json={'board':'chinext','company':company,'collection':collection,'filename':'公告.docx','file_base64':base64.b64encode(raw or docx_bytes()).decode()})
    assert r.status_code==200,r.text
    return r.json()


def commit(c,p,company='300001',**extra):
    return c.post('/api/library/imports/'+p['id']+'/commit',json={'board':'chinext','collection':p['collection'],'company':company,'metadata':{
        'title':'第三届董事会第二次会议决议公告','published_at':'2025-06-04','url':'https://www.cninfo.com.cn/test.pdf',
        'publication_confirmed':True,'company_confirmed':True,**extra}})


def test_history_upload_schedule_detail_and_company_isolation(workspace):
    c,root,_=workspace;p=upload(c);result=commit(c,p);assert result.status_code==200,result.text
    identity=result.json()['id'];s=c.get('/api/announcement-schedule?board=chinext&company=300001').json()
    assert len(s['items'])==1 and s['items'][0]['meeting']['term']==3 and s['items'][0]['meeting']['sequence']==2
    assert s['items'][0]['meeting_date']=='2025-06-03'
    assert c.get('/api/announcement-schedule?board=chinext&company=300002').json()['items']==[]
    assert c.get(f'/api/announcements/{identity}?board=chinext&company=300002').status_code==404
    detail=c.get(f'/api/announcements/{identity}?board=chinext&company=300001').json()
    assert detail['pages'][0]['text'].endswith('文末正文不可丢失。')
    assert any(b['type']=='table' for b in json.loads((root/detail['document_path']).read_text())['pages'][0]['blocks'])
    assert (root/detail['original_path']).read_bytes()==(root/p['original_path']).read_bytes()
    assert commit(c,p).status_code==200
    assert len(history.visible(history.state(root,'chinext','300001')))==1


def test_schedule_tag_filter_and_human_correction(workspace):
    c,root,_=workspace
    identity=commit(c,upload(c)).json()['id']
    path='/api/announcement-schedule?board=chinext&company=300001'
    result=c.get(path+'&tag=board&form=resolution').json()
    assert [r['id'] for r in result['items']]==[identity]
    assert c.get(path+'&tag=shareholder').json()['items']==[]
    assert c.get(path+'&tag=unknown').status_code==422
    body={'board':'chinext','company':'300001','id':identity,'fingerprint':result['fingerprint'],
          'changes':{'classification_override':{'tags':['shareholder'],'forms':['notice']}},'reason':'人工核对文件性质'}
    assert c.patch('/api/announcement-schedule',json={**body,'fingerprint':'stale'}).status_code==409
    assert c.patch('/api/announcement-schedule',json=body).status_code==200
    updated=c.get(path+'&tag=shareholder').json()
    assert updated['items'][0]['classification']['status']=='manual'
    assert c.get(path+'&tag=board').json()['total']==0
    row=history.state(root,'chinext','300001')['items'][0]
    assert row['corrections'][-1]['reason']==body['reason']
    assert c.patch('/api/announcement-schedule',json={**body,'fingerprint':updated['fingerprint'],'changes':{'classification_override':None}}).status_code==200
    assert c.get(path+'&tag=board').json()['total']==1
    assert c.get('/api/announcement-schedule?board=chinext&company=300002&tag=board').json()['total']==0


def test_web_import_automatically_starts_pi_classification(workspace,monkeypatch):
    from test_announcement_review import wait
    monkeypatch.setenv('DISCLOSURE_CLASSIFICATION_TEST','not-real-key')
    c,root,app=workspace;rt=app.state.pi_runtime
    rt.config_override={'models':[{'key':'test','label':'Test','id':'test-model','provider':'fixture','api':'openai-completions','enabled':True,'baseUrl':'http://127.0.0.1:1','api_key_env':'DISCLOSURE_CLASSIFICATION_TEST','contextWindow':200000,'maxTokens':4096}],'routes':{'default':'test'}}
    def classify(packet,emit,bridge,stop):
        d=json.loads(packet['system'].split('\n',1)[1])['documents'][0]
        bridge('submit_candidate',{'result':{'items':[{'id':d['id'],'status':'classified','tags':['board'],'forms':['resolution'],'page':1,'quote':d['title'],'reason':'标题为董事会会议决议'}]}})
        emit({'type':'done'})
    rt.runner=classify
    result=commit(c,upload(c)).json()
    assert result['classification_run']['status']=='accepted'
    assert wait(rt,result['classification_run']['run_id'])['status']=='completed'
    row=c.get('/api/announcement-schedule?board=chinext&company=300001').json()['items'][0]
    assert row['classification']['status']=='pi_reviewed'
    rt.close()


def test_unpublished_wrong_company_and_cross_scope_import_are_blocked(workspace):
    c,_,_=workspace;p=upload(c)
    assert commit(c,p,publication_confirmed=False).status_code==422
    assert commit(c,p,company='300002').status_code==403
    wrong=upload(c,docx_bytes('300002'))
    assert commit(c,wrong).status_code==422
    assert c.get(f"/api/library/imports/{p['id']}?board=chinext&company=300002").status_code==403
    assert not c.get('/api/announcement-schedule?board=chinext&company=300001').json()['items']


def test_cases_duplicate_does_not_replace_registered_parse(workspace,monkeypatch):
    c,root,_=workspace;raw='<html><p>官方案例原件。</p></html>'.encode()
    metadata={'title':'重复案例','url':'https://www.szse.cn/case-fixture','published_at':'2026-09-01',
              'stock_code':'300001','company_name':'测试公司300001','kind':'related_transaction'}
    first=file_ingestion.prepare(root,'chinext','cases',raw,'case.html',url=metadata['url'])
    assert file_ingestion.commit(root,first['id'],'chinext','cases','',metadata,'test')['status']=='updated'
    row=next(x for x in library_admin.state(root,'cases','chinext')['items'] if x['sha256']==library_admin.sha(raw))
    document_path=root/row['document_path'];before=document_path.read_bytes()
    changed=copy.deepcopy(first['document']);changed['pages'][0]['text']+=' 新解析版本。'
    monkeypatch.setattr(file_ingestion,'extract',lambda path:changed)
    second=file_ingestion.prepare(root,'chinext','cases',raw,'case.html',url=metadata['url'])
    result=file_ingestion.commit(root,second['id'],'chinext','cases','',metadata,'test')
    assert result['status']=='already_present' and result['id']==row['id']
    assert document_path.read_bytes()==before
    current=next(x for x in library_admin.state(root,'cases','chinext')['items'] if x['sha256']==library_admin.sha(raw))
    assert current['document_sha256']==library_admin.sha(before)


def test_history_soft_delete_reimport_keeps_old_parse_version(workspace,monkeypatch):
    c,root,_=workspace;raw=docx_bytes();metadata={'title':'第三届董事会第二次会议决议公告','published_at':'2025-06-04',
        'url':'https://www.cninfo.com.cn/test.pdf','publication_confirmed':True,'company_confirmed':True}
    first=file_ingestion.prepare(root,'chinext','history',raw,'公告.docx',company='300001',url=metadata['url'])
    applied=file_ingestion.commit(root,first['id'],'chinext','history','300001',metadata,'test')
    old=history.visible(history.state(root,'chinext','300001'))[0];old_document=(root/old['document_path']).read_bytes()
    before_fingerprint=history.state(root,'chinext','300001')['fingerprint']
    changed=copy.deepcopy(first['document']);changed['pages'][0]['text']+=' 新解析版本。'
    monkeypatch.setattr(file_ingestion,'extract',lambda path:changed)
    active_duplicate=file_ingestion.prepare(root,'chinext','history',raw,'公告.docx',company='300001',url=metadata['url'])
    active_result=file_ingestion.commit(root,active_duplicate['id'],'chinext','history','300001',metadata,'test')
    assert active_result['status']=='already_present' and history.state(root,'chinext','300001')['fingerprint']==before_fingerprint
    assert (root/old['document_path']).read_bytes()==old_document
    preview=c.post('/api/library/deletion-preview',json={'board':'chinext','company':'300001','collection':'history','ids':[applied['id']]}).json()
    assert c.post('/api/library/delete',json={'preview':preview,'fingerprint':preview['fingerprint']}).status_code==200
    second=file_ingestion.prepare(root,'chinext','history',raw,'公告.docx',company='300001',url=metadata['url'])
    restored=file_ingestion.commit(root,second['id'],'chinext','history','300001',metadata,'test')
    current=history.visible(history.state(root,'chinext','300001'))[0]
    assert restored['status']=='applied' and current['id']==old['id']
    assert current['document_path']!=old['document_path'] and (root/old['document_path']).read_bytes()==old_document
    assert '新解析版本' in (root/current['document_path']).read_text()


def test_ingestion_stale_or_conflicting_publish_never_overwrites_existing_bytes(workspace):
    c,root,_=workspace;raw='<html><p>原件一。</p></html>'.encode();raw2='<html><p>原件二。</p></html>'.encode()
    metadata={'title':'案例一','url':'https://www.szse.cn/case-fixture','published_at':'2026-09-01','stock_code':'300001','company_name':'测试公司300001','kind':'related_transaction'}
    first=file_ingestion.prepare(root,'chinext','cases',raw,'case.html',url=metadata['url'])
    file_ingestion.commit(root,first['id'],'chinext','cases','',metadata,'test')
    before_catalog=(root/'data/public/boards/chinext/catalog.json').read_bytes()
    second=file_ingestion.prepare(root,'chinext','cases',raw2,'case2.html',url=metadata['url'])
    with pytest.raises(HTTPException):file_ingestion.commit(root,second['id'],'chinext','cases','',metadata,'test',expected='stale')
    assert (root/'data/public/boards/chinext/catalog.json').read_bytes()==before_catalog
    target=root/'data/public/originals'/(library_admin.sha(raw2)+'.html');target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(b'corrupt')
    with pytest.raises(HTTPException):file_ingestion.commit(root,second['id'],'chinext','cases','',metadata,'test')
    assert target.read_bytes()==b'corrupt'
    target.unlink();outside=root/'outside-original.html';outside.write_bytes(b'preserve')
    target.symlink_to(outside)
    with pytest.raises(HTTPException):file_ingestion.commit(root,second['id'],'chinext','cases','',metadata,'test')
    assert outside.read_bytes()==b'preserve'


def test_delete_is_preview_bound_and_recoverable(workspace):
    c,root,_=workspace;identity=commit(c,upload(c)).json()['id']
    preview=c.post('/api/library/deletion-preview',json={'board':'chinext','company':'300001','collection':'history','ids':[identity]}).json()
    assert c.post('/api/library/delete',json={'preview':preview,'fingerprint':'wrong'}).status_code==409
    assert c.post('/api/library/delete',json={'preview':preview,'fingerprint':preview['fingerprint']}).status_code==200
    state=history.state(root,'chinext','300001');assert not history.visible(state)
    assert c.get(f'/api/announcements/{identity}?board=chinext&company=300001').status_code==404
    restored=c.post('/api/announcement-schedule/restore',json={'board':'chinext','company':'300001','id':identity,'fingerprint':state['fingerprint']})
    assert restored.status_code==200 and len(restored.json()['items'])==1
    preview=c.post('/api/library/deletion-preview',json={'board':'chinext','company':'300001','collection':'history','ids':[identity]}).json()
    assert c.post('/api/library/delete',json={'preview':preview,'fingerprint':preview['fingerprint']}).status_code==200
    raw=(root/history.state(root,'chinext','300001')['items'][0]['original_path']).read_bytes()
    assert commit(c,upload(c,raw)).status_code==200
    state=history.state(root,'chinext','300001')
    assert len(state['items'])==1 and len(history.visible(state))==1


def test_meeting_scope_postponement_and_concurrent_draft_conflicts():
    old={'id':'a','title':'第三届董事会第二次会议通知','meeting_date':'2025-06-03',**history.extract_fields('', '第三届董事会第二次会议通知')}
    old['meeting_date']='2025-06-03'
    event={'id':'new','title':'新公告','facts':{},'draft':{'documents':[{'title':'第三届董事会第二次会议决议公告','text':'本次会议于2025年6月4日召开。'}]}}
    snapshot={'items':[old],'coverage':{'complete':True}}
    assert history.check(event,snapshot)[0][0]['code']=='meeting_date_conflict'
    event['draft']['documents'][0]['title']='第四届董事会第二次会议决议公告'
    assert not history.check(event,snapshot)[0]
    event['draft']['documents'][0]['title']='关于第三届董事会第二次会议延期的公告';event['facts']['historical_links']=['a']
    assert not history.check(event,snapshot)[0]
    sibling={'id':'other','title':'并行事项','draft':{'text':'公告编号：2025-004','documents':[]}}
    event['draft']['documents'][0]['text']='公告编号：2025-004'
    assert any(e['code']=='concurrent_announcement_conflict' for e in history.check(event,{'items':[]},[sibling])[0])
    assert history.meeting_key(history.extract_fields('', '2025年第三次临时股东大会通知'))!=history.meeting_key(history.extract_fields('', '2025年年度股东大会通知'))


def test_history_version_invalidates_only_this_company_and_checks_bytes(workspace):
    c,root,_=workspace
    a={'id':'a','layer':'chinext','stock_code':'300001','company_id':'a','company_name':'测试公司300001','kind':'unclassified','facts':{},'title':'事项'}
    b={**a,'id':'b','stock_code':'300002','company_id':'b'}
    before_a=fingerprint(a,Seeds(root).for_event(a).catalog(),'assessment');before_b=fingerprint(b,Seeds(root).for_event(b).catalog(),'assessment')
    identity=commit(c,upload(c)).json()['id']
    assert before_a!=fingerprint(a,Seeds(root).for_event(a).catalog(),'assessment')
    assert before_b==fingerprint(b,Seeds(root).for_event(b).catalog(),'assessment')
    row=history.visible(history.state(root,'chinext','300001'))[0];(root/row['document_path']).write_text('tampered')
    assert identity in history.snapshot(root,a)['integrity_errors']


def test_ocr_is_targeted_and_missing_text_cannot_be_admitted(workspace,monkeypatch):
    from pypdf import PdfWriter
    c,root,_=workspace;writer=PdfWriter();writer.add_blank_page(width=100,height=100);out=io.BytesIO();writer.write(out)
    monkeypatch.setattr(file_ingestion,'ocr',lambda p:{'text':'','engine':'unavailable','review_required':True})
    p=upload(c,out.getvalue())
    assert p['document']['requires_ocr']
    assert commit(c,p,extraction_confirmed=True).status_code==409


def test_manual_template_replacement_preserves_bytes_and_rejects_stale_version(workspace):
    c,root,_=workspace;folder=root/'templates/boards/chinext';template=Document()
    for text in ('{{ title }}','{{ company_name }}','{{ status_label }}','{{ body_text }}'):template.add_paragraph(text)
    template.save(folder/'base.docx')
    profile={'id':'profile-base','title':'模板','kind':'board_resolution','library_board':'chinext','layers':['chinext'],'layout_profile_id':'layout','sections':[]}
    (root/'data/public/boards/chinext/profiles.json').write_text(json.dumps([profile]))
    entry={'id':'tpl-base','name':'模板','kind':'board_resolution','library_board':'chinext','profile_id':'profile-base','file':'boards/chinext/base.docx','layout_profile_id':'layout'}
    (folder/'manifest.json').write_text(json.dumps([entry]))
    template.sections[0].header.paragraphs[0].text='人工页眉';template.add_paragraph('用户保留文字')
    data=io.BytesIO();template.save(data);p=upload(c,data.getvalue(),collection='profiles')
    args={'board':'chinext','company':'300001','profile_id':'profile-base','import_id':p['id']}
    preview=c.post('/api/library/template-replacement',json=args).json()
    assert c.post('/api/library/template-replacement',json={**args,'apply':True,'fingerprint':'stale'}).status_code==409
    result=c.post('/api/library/template-replacement',json={**args,'apply':True,'fingerprint':preview['fingerprint']})
    assert result.status_code==200,result.text
    assert (root/'templates'/result.json()['file']).read_bytes()==data.getvalue()
    a=Seeds(root).for_event({'layer':'chinext','stock_code':'300001'}).templates()[0]
    b=Seeds(root).for_event({'layer':'chinext','stock_code':'300002'}).templates()[0]
    assert a['file']!=b['file'] and b['file']==entry['file']
    from scripts.word_renderer import render
    event={'company_name':'测试公司','title':'公告','plan':{'items':[]},'draft':{'text':'# 公告\n正式正文\n|项目|数值|\n|---|---|\n|甲|10|\n用户保留文字'}}
    rendered,_=render(event,{'template_authority':'user_uploaded'},root/'templates'/a['file'])
    doc=Document(io.BytesIO(rendered))
    assert doc.sections[0].header.paragraphs[0].text=='人工页眉'
    assert any(p.text=='用户保留文字' for p in doc.paragraphs)
    assert doc.tables[0].cell(1,1).text=='10'
    assert (root/'templates'/a['file']).read_bytes()==data.getvalue()
    from backend.artifacts import check
    verified=check(rendered,event,{'template_authority':'user_uploaded'},data.getvalue())
    assert verified['status']=='PASS',verified
    doc.tables[0].cell(1,1).text='999';changed=io.BytesIO();doc.save(changed)
    assert check(changed.getvalue(),event,{'template_authority':'user_uploaded'},data.getvalue())['status']=='FAIL'


def test_new_uploaded_template_is_scoped_to_company(workspace):
    c,root,_=workspace;folder=root/'templates/boards/chinext';doc=Document()
    for value in ('{{ title }}','{{ company_name }}','{{ status_label }}','{{ body_text }}'):doc.add_paragraph(value)
    doc.save(folder/'base.docx')
    profile={'id':'p-base','title':'基础文种','kind':'board_resolution','library_board':'chinext','layers':['chinext'],'layout_profile_id':'layout','sections':[{'id':'body','title':'正文','fields':[{'prompt':'内容'}]}],'normative_source_ids':['law'],'case_evidence':[],'scope_review_status':'format_and_case_evidence_pending','official_format_number':'test'}
    public=root/'data/public/boards/chinext'
    (public/'profiles.json').write_text(json.dumps([profile]));(public/'catalog.json').write_text(json.dumps({'board':'chinext','sources':[{'id':'law','title':'测试法源','library_board':'chinext'}],'cases':[]}))
    (folder/'layout_profiles.json').write_text(json.dumps([{'id':'layout','library_board':'chinext'}]))
    (folder/'manifest.json').write_text(json.dumps([{'id':'tpl-base','name':'基础文种','kind':'board_resolution','library_board':'chinext','profile_id':'p-base','file':'boards/chinext/base.docx','layout_profile_id':'layout'}]))
    out=io.BytesIO();doc.save(out);p=upload(c,out.getvalue(),collection='profiles')
    result=c.post('/api/library/template-upload',json={'board':'chinext','company':'300001','base_profile_id':'p-base','title':'本公司手工模板','import_id':p['id']})
    assert result.status_code==200,result.text
    a=c.get('/api/library/search?board=chinext&company=300001&collection=profiles').json()['items']
    b=c.get('/api/library/search?board=chinext&company=300002&collection=profiles').json()['items']
    assert any(r['title']=='本公司手工模板' for r in a)
    assert all(r['title']!='本公司手工模板' for r in b)


def test_pi_uses_the_same_import_service_and_cannot_sign_publication(workspace):
    from backend import knowledge_ops
    from uuid import uuid4
    c,root,app=workspace;p=upload(c);runtime=app.state.pi_runtime
    session=runtime.store.create_session('chinext','','文件核对',str(uuid4()))
    run,_=runtime.store.accept(session,{'stage':'auto','text':'入库文件','model_key':'fixture','company_code':'300001','request_id':str(uuid4())},{})
    runtime.store.update(run['id'],stage='knowledge',intent_domain='knowledge',intent='refresh')
    source=knowledge_ops.execute(runtime,run['id'],'knowledge_import_read',{'import_id':p['id']})
    assert '文末正文不可丢失' in source['data']['page']['text']
    knowledge_ops.execute(runtime,run['id'],'knowledge_import_propose',{'import_id':p['id'],'metadata':{'title':'第三届董事会第二次会议决议公告','url':'https://www.cninfo.com.cn/test.pdf','published_at':'2025-06-04','publication_confirmed':True},'summary':'登记公告'})
    change=runtime.store.run(run['id'])['knowledge_change'];assert 'publication_confirmed' not in change['metadata']
    runtime.store.update(run['id'],status='waiting_knowledge_confirmation')
    with pytest.raises(HTTPException):knowledge_ops.confirm(runtime,run['id'],change['fingerprint'],True)
    assert not history.visible(history.state(root,'chinext','300001'))
    knowledge_ops.confirm(runtime,run['id'],change['fingerprint'],True,{'publication_confirmed':True,'company_confirmed':True})
    assert len(history.visible(history.state(root,'chinext','300001')))==1
    assert runtime.store.run(run['id'])['status']=='completed'


def test_pi_batch_resumes_without_duplicate_imports(workspace,monkeypatch):
    from backend import knowledge_ops
    from uuid import uuid4
    c,root,app=workspace;p=upload(c);p2=upload(c,docx_bytes(title='第三届董事会第三次会议决议公告',announcement_number='2025-003'))
    runtime=app.state.pi_runtime;s=runtime.store.create_session('chinext','','批量测试',str(uuid4()))
    run,_=runtime.store.accept(s,{'stage':'auto','text':'批量入库','model_key':'fixture','company_code':'300001','request_id':str(uuid4())},{})
    runtime.store.update(run['id'],stage='knowledge',intent_domain='knowledge',intent='refresh')
    items=[{'import_id':x['id'],'metadata':{'title':x['suggested_title'],'published_at':'2025-06-04','url':'https://www.cninfo.com.cn/test.pdf'}} for x in (p,p2)]
    knowledge_ops.execute(runtime,run['id'],'knowledge_import_batch_propose',{'imports':items,'summary':'两份文件一次确认'})
    change=runtime.store.run(run['id'])['knowledge_change'];runtime.store.update(run['id'],status='waiting_knowledge_confirmation')
    actual=file_ingestion.commit
    def fail_second(*args,**kwargs):
        if args[1]==p2['id']:raise HTTPException(409,'隔离模拟中断')
        return actual(*args,**kwargs)
    monkeypatch.setattr(file_ingestion,'commit',fail_second)
    flags={'publication_confirmed':True,'company_confirmed':True}
    with pytest.raises(HTTPException):knowledge_ops.confirm(runtime,run['id'],change['fingerprint'],True,flags)
    assert len(history.visible(history.state(root,'chinext','300001')))==1
    assert runtime.store.run(run['id'])['knowledge_change']['status']=='partial'
    monkeypatch.setattr(file_ingestion,'commit',actual)
    knowledge_ops.confirm(runtime,run['id'],change['fingerprint'],True,flags)
    assert len(history.visible(history.state(root,'chinext','300001')))==2
    assert runtime.store.run(run['id'])['knowledge_change']['result']['remaining']==0


def test_pdf_structure_marks_table_candidates_for_review():
    blocks=file_ingestion.pdf_blocks('一、基本情况\n项目    金额\n收入    100\n\n后续段落',1)
    assert blocks[0]['type']=='heading'
    table=next(b for b in blocks if b['type']=='table_candidate')
    assert table['rows']==[['项目','金额'],['收入','100']] and table['review_required']


def test_company_sqlite_has_document_and_internal_unit_indexes(workspace):
    from backend import announcement_index
    c,root,_=workspace;commit(c,upload(c))
    status=announcement_index.status(root,'chinext','300001')
    assert status['status']=='current' and status['announcements']==1 and status['pages']==1
    assert status['units']>=4 and not status['integrity_errors']
    assert status['coverage']=={'active_catalog_items':1,'primary_indexed':1,'page_records':1,'expected_pages':1,'secondary_indexed_documents':1,'large_documents':0,'large_documents_indexed':0}
    with announcement_index.connect(announcement_index.paths(root,'chinext','300001')[2],True) as db:
        assert db.execute('SELECT COUNT(*) FROM document_pages').fetchone()[0]==1
    result=announcement_index.search(root,'chinext','300001','经营计划')
    assert result['items'][0]['title']=='第三届董事会第二次会议决议公告'
    passage=result['passages'][0]
    assert passage['page']==1 and passage['unit_type']=='table' and passage['anchor'].startswith('body:')
    routed=history.search(root,{'layer':'chinext','stock_code':'300001'},'经营计划')
    assert routed['retrieval_backend']=='company_sqlite_fts5' and routed['catalog_indexed']
    assert routed['passages'][0]['page']==1


def test_history_commit_reports_index_failure_instead_of_silent_success(workspace,monkeypatch):
    from backend import announcement_index
    c,root,_=workspace;p=upload(c);actual=announcement_index.sync
    monkeypatch.setattr(announcement_index,'sync',lambda *args:(_ for _ in ()).throw(RuntimeError('index failed')))
    result=commit(c,p)
    assert result.status_code==200 and result.json()['index']['status']=='stale'
    assert len(history.visible(history.state(root,'chinext','300001')))==1
    monkeypatch.setattr(announcement_index,'sync',actual)
    retried=commit(c,p)
    assert retried.status_code==200 and retried.json()['index']['status']=='current'
    assert len(history.visible(history.state(root,'chinext','300001')))==1


def test_pi_marks_index_failure_partial_and_retry_only_rebuilds(workspace,monkeypatch):
    from backend import announcement_index,knowledge_ops
    from uuid import uuid4
    c,root,app=workspace;p=upload(c);runtime=app.state.pi_runtime
    session=runtime.store.create_session('chinext','','索引失败',str(uuid4()))
    run,_=runtime.store.accept(session,{'stage':'auto','text':'入库','model_key':'fixture','company_code':'300001','request_id':str(uuid4())},{})
    runtime.store.update(run['id'],stage='knowledge',intent_domain='knowledge',intent='refresh',status='waiting_knowledge_confirmation')
    knowledge_ops.execute(runtime,run['id'],'knowledge_import_propose',{'import_id':p['id'],'metadata':{'title':'第三届董事会第二次会议决议公告','url':'https://www.cninfo.com.cn/test.pdf','published_at':'2025-06-04'},'summary':'入库'})
    change=runtime.store.run(run['id'])['knowledge_change'];actual=announcement_index.sync
    monkeypatch.setattr(announcement_index,'sync',lambda *args:(_ for _ in ()).throw(RuntimeError('index failed')))
    with pytest.raises(HTTPException):knowledge_ops.confirm(runtime,run['id'],change['fingerprint'],True,{'publication_confirmed':True,'company_confirmed':True})
    assert runtime.store.run(run['id'])['knowledge_change']['status']=='partial'
    monkeypatch.setattr(announcement_index,'sync',actual)
    result=knowledge_ops.confirm(runtime,run['id'],change['fingerprint'],True,{'publication_confirmed':True,'company_confirmed':True})
    assert result['status']=='applied' and result['result']['index']['status']=='current'
    assert len(history.visible(history.state(root,'chinext','300001')))==1
