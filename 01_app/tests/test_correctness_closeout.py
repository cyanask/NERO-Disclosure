"""Small regressions for evidence, lossless draft export and live-data isolation."""
import hashlib
import io
import sqlite3
from pathlib import Path

import pytest
from docx import Document
from fastapi import HTTPException
from backend import consult_export as export, disclosure_contract as contract
from backend.document_content import WordDeliveryError
from conftest import KNOWLEDGE, LOCAL, ROOT, seed_tree
from test_consult_export import runtime, settled_run, LAYOUT, LAW, assessment
from test_semantic_review import event, default_matter, result_with_matter, fact
from test_agent_tasks import env, claimed, candidate, post, latest
from test_open_intake import real_event


@pytest.mark.parametrize('operation', ['write', 'unlink', 'sqlite'])
def test_live_data_is_protected_before_write(operation):
    path = KNOWLEDGE / 'data/public/boards/chinext/catalog.json'
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    with pytest.raises(PermissionError, match='测试禁止写入正式资料'):
        if operation == 'write':
            path.write_text('must never be written')
        elif operation == 'unlink':
            path.unlink()
        else:
            sqlite3.connect(LOCAL / 'var/disclosure.sqlite3')
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_seed_copy_is_not_a_hardlink_or_live_root(tmp_path):
    source = tmp_path / 'source'; source.mkdir()
    (source / 'a.txt').write_text('original')
    target = seed_tree(source, tmp_path / 'copy')
    (target / 'a.txt').write_text('changed in test')
    assert (source / 'a.txt').read_text() == 'original'
    with pytest.raises(ValueError, match='项目外'):
        seed_tree(source, KNOWLEDGE / 'data/forbidden-test-copy')


def negative_result():
    matter = default_matter()
    matter.update(duty_status='not_triggered', reassessment_conditions=['事实变化后复判'],
                  reasoning_items=[{'source_id':'law-1','locator':'第一条','quote':LAW['text'],
                                    'fact_keys':['amount'],'outcome':'not_met'}])
    value = result_with_matter([fact(value=100, status='material_supported',
                                   source_ref='announcement-missing', quote='')], matter)
    value.update(status='no_disclosure', source_ids=['law-1'])
    return value


def test_no_disclosure_cannot_use_unlocatable_material():
    value = negative_result()
    errors, _, _ = contract.assessment_checks(event(), {'sources':[LAW]}, value, history_reader=lambda _: {})
    assert 'history_fact_unlocatable' in {e['code'] for e in errors}


def test_calculation_cannot_use_unlocatable_material():
    value = negative_result()
    value['matters'][0]['reasoning_items'] = []
    value['matters'][0]['calculations'] = [{'operation':'sum','scope':'snapshot','unit':'元',
        'fact_keys':['amount'],'basis_source_id':'facts','result':'100','aggregation_basis':'public_announcements'}]
    errors, _, _ = contract.assessment_checks(event(), {'sources':[LAW]}, value, history_reader=lambda _: {})
    assert 'history_fact_unlocatable' in {e['code'] for e in errors}


def test_ordinary_gap_continues_but_is_not_verified_material():
    value = negative_result()
    value['matters'][0]['reasoning_items'] = []
    errors, warnings, _ = contract.assessment_checks(event(), {'sources':[LAW]}, value, history_reader=lambda _: {})
    assert 'history_fact_unlocatable' not in {e['code'] for e in errors}
    assert 'fact_pending_supplement' in {e['code'] for e in warnings}
    assert value['facts'][0]['status'] == 'unknown'
    assert value['missing']


def test_admitted_gap_is_pending_in_saved_event_and_displayed_result(env):
    client, tokens, _, _ = env
    case = real_event(client)
    task, lease = claimed(client, case, tokens[0])
    value = candidate()
    value['assessment_as_of'] = case['facts']['assessment_as_of']
    for row in value['facts']:
        row['observed_at'] = value['assessment_as_of']
    value['facts'].append({'key':'ordinary_gap','value':'待核背景','status':'material_supported',
                          'source_ref':'announcement-missing','quote':'','observed_at':value['assessment_as_of']})
    response = post(client,case,f"agent-tasks/{task['id']}/evaluate",tokens[0],
                    claim_id=lease['claim_id'],input_fingerprint=lease['input_fingerprint'],result=value)
    assert response.status_code == 200, response.text
    assert response.json()['outcome'] == 'waiting_approval', response.json().get('gate')
    for facts in (latest(client,case)['assessment']['facts'],response.json()['result_snapshot']['result']['facts']):
        assert next(f for f in facts if f['key']=='ordinary_gap')['status'] == 'unknown'


def test_explicit_bad_quote_is_not_rebound_to_matching_value():
    value = negative_result()
    value['facts'][0].update(status='user_statement', source_ref='summary', value='重大事项', quote='捏造的原文')
    errors, _, _ = contract.assessment_checks(event(summary='公司发生重大事项。'), {'sources':[LAW]}, value)
    assert 'fact_source_mismatch' in {e['code'] for e in errors}
    assert value['facts'][0]['quote'] == '捏造的原文'
    assert contract.bind_summary_quote('预算100元；实际支出100元。', 100) is None
    assert contract.bind_summary_quote('公司发生重大事项；另一公司发生重大事项。','重大事项') is None


def body_text(raw):
    doc = Document(io.BytesIO(raw))
    return '\n'.join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for r in t.rows for c in r.cells])


def test_word_keeps_ragged_table_rows_links_numbers_and_gaps():
    text = ('# 测试咨询\n3. 对应附件三\n4. 对应附件四\n'
            '依据：[交易所原文](https://example.com/official)\n'
            '| 事项 | 金额 |\n| --- | --- |\n| 甲 | 100 |\n| 乙 |\n| 丙 | 300 | 【待补】 |')
    actual = body_text(export.render(text,'测试咨询','测试公司',LAYOUT,'2026-09-16'))
    for required in ['3. 对应附件三','4. 对应附件四','https://example.com/official',
                     '| 甲 | 100 |','| 乙 |','| 丙 | 300 | 【待补】 |']:
        assert required in actual


def test_word_readback_rejects_missing_content():
    doc = Document();doc.add_paragraph('only part of the reply')
    out = io.BytesIO();doc.save(out)
    with pytest.raises(WordDeliveryError, match='回读与已确认正文不一致'):
        export.verify_rendered_content(out.getvalue(),[],[{'kind':'paragraph','text':'complete approved reply'}])


@pytest.mark.parametrize('original,changed', [('甲\n\n乙','甲乙'), ('金额 1 000 元','金额 1000 元')])
def test_export_binds_paragraphs_and_interior_spaces(tmp_path, original, changed):
    value = runtime(tmp_path,settled_run(),[{'kind':'assistant','body':{'phase':'final','text':original}}])
    with pytest.raises(HTTPException) as exc:
        export.create(value,'sid-1',{'run_id':'run-1','text':changed})
    assert exc.value.status_code == 410
    assert not (tmp_path / 'work/consult-exports').exists()


def test_export_refuses_progress_or_formal_workflow_bypass(tmp_path):
    value = runtime(tmp_path,settled_run(),[{'kind':'assistant','body':{'phase':'progress','text':'working'}}])
    with pytest.raises(HTTPException):export.create(value,'sid-1',{'run_id':'run-1','text':'working'})
    value = runtime(tmp_path,{**settled_run(),'stage':'draft'},[{'kind':'assistant','body':{'text':'draft'}}])
    with pytest.raises(HTTPException):export.create(value,'sid-1',{'run_id':'run-1','text':'draft'})
