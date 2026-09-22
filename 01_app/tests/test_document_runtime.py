"""Real Pi routing, Word subprocess and local APIs; only the provider is simulated."""
import io
import json
import threading
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile
import pytest
from docx import Document
from backend import document_store, document_runtime, document_files
from test_pi_runtime import client, new_session, settled


def request(c, session, text='认可你的方案，请制作成 Word', request_id=None):
    event_id=c.app.state.pi_runtime.store.session(session['id'])['event_id']
    revision=c.get('/api/events/'+event_id).json()['revision'] if not event_id.startswith('conversation:') else 0
    return c.post(f'/api/chat/sessions/{session["id"]}/runs', json={
        'text': text, 'model_key': 'fixture-a', 'expected_revision': revision, 'request_id': request_id or str(uuid4())})


def readiness(document, gaps=()):
    """Drafting-readiness entry for one document; identity follows its saved id."""
    entry = {'title': document['title'], 'kind': document['kind'],
             'checked_topics': ['主体、事项、时间和金额'], 'gaps': [dict(gap) for gap in gaps]}
    if document.get('document_id'):
        entry['document_id'] = document['document_id']
    return entry


def content_gap(label_text, reason='资料尚未提供', state='missing'):
    return {'label': label_text, 'reason': reason, 'category': 'content', 'state': state}


def reviewed(run):
    """Support consultation review in shared fixtures; document saves do not review."""
    def respond(packet, emit, bridge, stop):
        if packet.get('stage')=='semantic_review':
            items=json.loads(packet['prompt'].split('\n',1)[1])['items']
            verdicts=[{'item_id':row['item_id'],'verdict':'supported','reason':'离线制文回归用例的已知支持结论'} for row in items]
            emit({'type':'assistant','message':0,'phase':'answer','stopReason':'stop','text':json.dumps({'verdicts':verdicts})})
            emit({'type':'done'});return
        return run(packet,emit,bridge,stop)
    return respond


def provider(runtime, build, gaps=None, consent=None):
    """Tools are available by domain; preflight still validates the document before saving."""
    def run(packet, emit, bridge, stop):
        emit({'type': 'started'})
        routed = bridge('route_request', {'domain': 'disclosure', 'intent': 'document', 'reason': '用户要求制作Word'})
        names = [t['name'] for t in routed['next_context']['tools']]
        assert routed['next_context']['stage'] == 'document_preflight'
        assert 'assess_document_readiness' in names and 'make_word' in names
        context = bridge('read_document_context', {})['data']
        documents = build(context, bridge, stop)
        assessment = {'documents': [readiness(document, (gaps or {}).get(document['title'], [])) for document in documents],
                      'request_quote': packet['prompt'], 'decision': 'assess'}
        if consent:
            assessment.update(decision='proceed_with_gaps', notice_id=consent, choice_quote=packet['prompt'])
        opened = bridge('assess_document_readiness', assessment)
        assert 'next_context' in opened, opened
        assert 'make_word' in [t['name'] for t in opened['next_context']['tools']], opened
        result = bridge('make_word', {'documents': documents})
        assert result['data']['business_state_changed'] is False
        emit({'type': 'assistant', 'phase': 'final', 'stopReason': 'stop', 'text': '文档已生成，待审阅。'})
        emit({'type': 'done'})
    runtime.runner = reviewed(run)


def draft(title='分析材料', **kwargs):
    return {'title': title, 'kind': 'analysis', 'template_id': 'builtin:analysis',
            'text': '# '+title+'\n\n一、当前结论\n以当前资料为准。\n\n二、后续安排\n请核对实施状态。',
            'pending': [], 'basis': [], **kwargs}


def listing(c, s):
    return c.get(f'/api/chat/sessions/{s["id"]}/documents').json()


def test_short_request_generates_from_consultation_without_business_event(client):
    c, runtime, _ = client
    s = runtime.store.create_session('chinext', '', '咨询', str(uuid4()))
    provider(runtime, lambda *_: [draft()])
    rid = str(uuid4())
    result = settled(c, request(c, s, request_id=rid))
    assert result['run']['status'] == 'completed', result['run']
    assert runtime.store.session(s['id'])['event_id'].startswith('conversation:')
    row = listing(c, s)['items'][0]
    raw = c.get(row['download']).content
    doc = Document(io.BytesIO(raw))
    assert [p.text for p in doc.paragraphs].count('分析材料') == 1
    assert row['review_status'] == 'pending'
    assert any(r['kind'] == 'script_started' for r in result['events'])
    repeat = request(c, s, request_id=rid)
    assert repeat.json()['id'] == result['run']['id']
    assert listing(c, s)['items'][0]['version'] == 1


def test_mid_workflow_announcement_batch_does_not_advance_or_confirm(client):
    c, runtime, _ = client
    s, event = new_session(c)
    before = c.get('/api/events/'+event['id']).json()
    def build(ctx, *_):
        selected = next(t for t in ctx['templates'] if t.get('kind') == 'related_transaction')
        source = next(s for s in ctx['sources'] if s['id'].startswith('user:'))
        return [draft('财务资助公告草稿', kind='announcement', template_id=selected['id'],
                      text='# 财务资助公告草稿\n一、事项概述\n拟提供财务资助1000万元。\n二、审议情况\n【待补：审议情况】',
                      pending=['审议情况'], basis=[{'source_id':source['id'], 'quote':'拟提供财务资助1000万元',
                                                    'statement':'拟提供财务资助1000万元。'}]), draft('事项分析')]
    gaps = {'财务资助公告草稿': [content_gap('审议情况', '董事会审议情况尚未提供')]}
    def notice(packet, emit, bridge, stop):
        bridge('route_request', {'domain': 'disclosure', 'intent': 'document', 'reason': '用户要求制作Word'})
        documents = build(bridge('read_document_context', {})['data'], bridge, stop)
        result = bridge('assess_document_readiness', {
            'documents': [readiness(document, gaps.get(document['title'], [])) for document in documents],
            'request_quote': packet['prompt'], 'decision': 'assess'})
        assert result['terminate'] and result['data']['status'] == 'waiting_user', result
        emit({'type': 'done'})
    runtime.runner = reviewed(notice)
    notice = settled(c, request(c, s, '拟提供财务资助1000万元。请将公告草稿和分析稿制作成Word'))
    assert notice['run']['status'] == 'waiting_user', notice
    assert listing(c, s)['items'] == []
    provider(runtime, build, gaps=gaps, consent=notice['run']['document_preflight']['notice_id'])
    result = settled(c, request(c, s, '先按现有资料制作Word，缺项标注待补'))
    assert result['run']['status'] == 'completed', result
    assert c.get('/api/events/'+event['id']).json() == before
    rows = listing(c, s)['items']
    assert len(rows) == 2 and '审议情况' in rows[0]['pending']
    assert '咨询工作稿' not in c.get(rows[0]['download']).text
    assert all(c.get(r['download']).status_code == 200 for r in rows)
    assert c.post(f'/api/chat/sessions/{s["id"]}/documents/{rows[0]["document_id"]}/review',
                  json={'version':1,'sha256':rows[0]['sha256'],'decision':'accepted','content_reviewed':True,'visual_reviewed':True}).status_code == 409


def test_no_file_tool_is_not_reported_as_completed(client):
    c, runtime, _ = client
    s, _ = new_session(c)
    def run(packet, emit, bridge, stop):
        bridge('route_request', {'domain':'disclosure','intent':'document','reason':'制作文档'})
        emit({'type':'assistant','text':'文件已完成'})
        emit({'type':'done'})
    runtime.runner = reviewed(run)
    result = settled(c, request(c, s))
    assert result['run']['status'] == 'incomplete'
    assert listing(c, s)['items'] == []


def test_reported_five_requests_keep_consent_and_deliver_real_word(client):
    c,runtime,_=client
    s=runtime.store.create_session('chinext','','董秘公告回归',str(uuid4()))
    texts=['请根据以下已知情况起草公告，缺失信息请具体标注待补：公司准备更换董秘',
           '选择B','要按照待补材料推进，先生成Word。',
           '直接给我先制造一个公告Word里面没有的资料先空着，我后面手工补。','就按照这样的内容制造Word。']
    def run(packet,emit,bridge,stop):
        word='Word' in packet['prompt']
        bridge('route_request',{'domain':'disclosure','intent':'document' if word else 'announcement','reason':'承接同一公告的待补要求'})
        ctx=bridge('read_document_context',{})['data']
        template=next(t for t in ctx['templates'] if t['kind']=='management_change')
        old=next(iter(ctx['documents']),None)
        document=draft('董事会秘书变动公告',kind='announcement',template_id=template['id'],
            text='# 董事会秘书变动公告\n一、人员变动情况概述\n【待补：离任及审议情况】\n二、新任人员简历及合规声明\n【待补：简历及任职合规事实】\n三、备查文件\n【待补：备查文件】')
        gaps=[content_gap(label) for label in ['离任及审议情况','简历及任职合规事实','备查文件']]
        if old:
            document.update(document_id=old['document_id'],base_version=old['version'])
            if old.get('format')=='docx':document['text']=''
        args={'documents':[readiness(document,gaps)],'request_quote':packet['prompt'],'decision':'assess'}
        # The model judges the request; the runtime only records the exact user message.
        if any(marker in packet['prompt'] for marker in ('待补','留空','空着')):
            args.update(decision='draft_with_placeholders',choice_quote=packet['prompt'])
        opened=bridge('assess_document_readiness',args)
        assert 'next_context' in opened,opened
        production={'documents':[document]}
        if not word:production['assessment']={'disclosure_needed':'uncertain','disclosure_scope':'董事会秘书变动','reason':'仅制作待补稿，具体事实尚未提供'}
        bridge('make_word' if word else 'save_announcement',production)
        emit({'type':'done'})
    runtime.runner = reviewed(run)
    for text in texts:
        result=settled(c,request(c,s,text))
        assert result['run']['status']=='completed',result
        assert not any(e['kind']=='document_gap_notice' for e in result['events'])
    rows=listing(c,s)['items']
    assert len(rows)==1 and rows[0]['format']=='docx'
    raw=c.get(rows[0]['download']).content
    doc=Document(io.BytesIO(raw))
    body='\n'.join(p.text for p in doc.paragraphs)
    assert '【待补：简历及任职合规事实】' in body
    assert '未受处罚' not in body and '不存在' not in body
    assert rows[0]['review_status']=='pending'


def test_document_revisions_and_hash_bound_review_leave_event_unchanged(client):
    c, runtime, _ = client
    s, e = new_session(c)
    provider(runtime, lambda *_: [draft()])
    settled(c, request(c, s))
    prior = listing(c, s)['items'][0]
    before = c.get('/api/events/'+e['id']).json()
    response = c.post(f'/api/chat/sessions/{s["id"]}/documents/{prior["document_id"]}/review',
                      json={'version':1,'sha256':prior['sha256'],'decision':'accepted','content_reviewed':True,'visual_reviewed':True})
    assert response.status_code == 200, response.text
    def revise(context, bridge, stop):
        # 已确认稿按锚点局部修订；整篇覆盖只适用于未确认的系统生成稿。
        read = bridge('read_document', {'document_id': prior['document_id']})['data']
        anchor = next(b for b in read['blocks'] if b['text'] == '以当前资料为准。')
        return [draft(document_id=prior['document_id'], base_version=1, template_id='source:current', text='',
                      edits=[{'block_id': anchor['id'], 'original': '以当前资料为准。', 'replacement': '修订后的内容。'}])]
    provider(runtime, revise)
    result = settled(c, request(c, s, '修改当前文稿，重新生成Word'))
    assert result['run']['status'] == 'completed', result
    current = listing(c, s)['items'][0]
    assert current['version'] == 2 and current['review_status'] == 'pending'
    assert current['history'][1]['review_status'] == 'accepted'
    assert c.get(prior['download']).status_code == 200
    assert c.get('/api/events/'+e['id']).json() == before
    assert c.post(f'/api/chat/sessions/{s["id"]}/documents/{prior["document_id"]}/review',
                  json={'version':1,'sha256':prior['sha256'],'decision':'accepted','content_reviewed':True,'visual_reviewed':True}).status_code == 409


def test_accepted_word_cannot_be_overwritten_wholesale(client):
    """已确认稿只能按锚点局部修订；整篇覆盖不得当作新版本保存。"""
    c,runtime,_=client;s,_=new_session(c)
    provider(runtime,lambda *_:[draft()]);settled(c,request(c,s));row=listing(c,s)['items'][0]
    assert c.post(f'/api/chat/sessions/{s["id"]}/documents/{row["document_id"]}/review',json={
        'version':1,'sha256':row['sha256'],'decision':'accepted','content_reviewed':True,'visual_reviewed':True}).status_code==200
    provider(runtime,lambda *_:[draft(document_id=row['document_id'],base_version=1,text='# 分析材料\n整篇覆盖后的内容。')])
    out=settled(c,request(c,s,'把已确认稿整篇重写并生成Word'))
    assert out['run']['status']=='failed'
    assert '锚点' in out['run'].get('reason','')
    assert listing(c,s)['items'][0]['version']==1


def legacy_manual_source(runtime, sid, raw, **kwargs):
    """Compatibility fixture for versions registered before return imports retired."""
    return document_store.publish(runtime, sid, [{
        'title':'人工修改稿', 'kind':'analysis', 'raw':raw,
        'source_type':'human_import', 'template_id':'source:current',
        'snapshot':document_files.inspect(raw), **kwargs,
    }], document_store.load(runtime, sid)['revision'])[0]


def test_legacy_manual_source_and_anchored_change_preserve_other_parts(client):
    c, runtime, _ = client
    s, e = new_session(c)
    document = Document()
    p = document.add_paragraph(); p.add_run('人工').bold=True;p.add_run('保留段落')
    document.add_paragraph('金额为1000万元。')
    document.add_table(rows=1, cols=1).cell(0,0).text='保持表格'
    stream=io.BytesIO();document.save(stream);raw=stream.getvalue()
    source=legacy_manual_source(runtime,s['id'],raw)
    def build(ctx,bridge,*_):
        read=bridge('read_document',{'document_id':source['document_id']})['data']
        anchor=next(b for b in read['blocks'] if b['text']=='金额为1000万元。')
        return [draft(title='人工修改稿',document_id=source['document_id'],base_version=1,template_id='source:current',text='',
                      edits=[{'block_id':anchor['id'],'original':'1000','replacement':'1200'}])]
    provider(runtime,build)
    result=settled(c,request(c,s,'将人工稿中的金额改为1200万元并制作Word'))
    assert result['run']['status']=='completed',result
    new=listing(c,s)['items'][0]
    assert new['source_type']=='manual_revision'
    changed=c.get(new['download']).content
    with ZipFile(io.BytesIO(raw)) as a,ZipFile(io.BytesIO(changed)) as b:
        assert a.namelist()==b.namelist()
        assert [n for n in a.namelist() if a.read(n)!=b.read(n)]==['word/document.xml']
    final=Document(io.BytesIO(changed))
    assert final.paragraphs[0].runs[0].bold and final.paragraphs[1].text=='金额为1200万元。'
    assert final.tables[0].cell(0,0).text=='保持表格'
    assert c.get(source['download']).content==raw


def test_cross_session_is_rejected_but_missing_basis_does_not_shorten_draft(client):
    c,runtime,_=client;s,_=new_session(c)
    provider(runtime,lambda *_:[draft()]);settled(c,request(c,s));row=listing(c,s)['items'][0]
    other=runtime.store.create_session('chinext','','另一会话',str(uuid4()))
    assert c.get(f'/api/chat/sessions/{other["id"]}/documents/{row["document_id"]}/versions/1/file').status_code==404
    text='会话问题及待核实的分析应完整保留。'
    basis=[{'statement':text,'source_id':'user:nonexistent','quote':'尚未定位的来源'}]
    def run(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','reason':'用户要求制作Word'})
        bridge('read_document_context',{})
        bridge('assess_document_readiness',{'documents':[readiness(draft())],'request_quote':packet['prompt'],'decision':'assess'})
        result=bridge('make_word',{'documents':[draft(text=text,basis=basis)]})
        assert result['data']['business_state_changed'] is False
        emit({'type':'done'})
    runtime.runner = reviewed(run)
    out=settled(c,request(c,other))
    assert out['run']['status']=='completed',out
    assert not any(e['kind']=='document_evidence_repair' for e in out['events'])
    saved=listing(c,other)['items'][0]
    assert saved['pending']==[]
    assert text in document_files.inspect(c.get(saved['download']).content)['text']
    snapshot=document_store.snapshot(runtime,other['id'],saved['document_id'])
    assert snapshot['basis']==basis
    assert saved['checks']['evidence_bindings']=='recorded_unverified'


def test_cancellation_before_registration_preserves_previous_version(client,monkeypatch):
    c,runtime,_=client;s,_=new_session(c)
    provider(runtime,lambda *_:[draft()]);settled(c,request(c,s));before=listing(c,s)
    original=document_runtime.run_script
    def stopped(runtime,rid,packet,stop):
        result=original(runtime,rid,packet,stop);stop.set();return result
    monkeypatch.setattr(document_runtime,'run_script',stopped)
    row=before['items'][0]
    provider(runtime,lambda *_:[draft(document_id=row['document_id'],base_version=1,text='新版本正文')])
    result=settled(c,request(c,s,'修改并生成Word'))
    assert result['run']['status']=='cancelled'
    assert listing(c,s)==before


def test_legacy_export_post_is_retired(client):
    c,runtime,_=client;s,_=new_session(c)
    assert c.post(f'/api/chat/sessions/{s["id"]}/exports',json={'run_id':'any','text':'不应生成'}).status_code==410
    assert c.get(f'/api/chat/sessions/{s["id"]}/exports').status_code==200
    assert listing(c,s)['items']==[]


@pytest.mark.parametrize('text_only',[True,False])
def test_incomplete_profit_and_investment_draft_is_saved_with_gaps(client,text_only):
    c,runtime,_=client;s,e=new_session(c)
    before=c.get('/api/events/'+e['id']).json()
    gaps=[content_gap('分红金额','每十股派发的现金红利金额尚未确定'),content_gap('董事会届次和表决结果','审议程序要素尚未提供')]
    def build(bridge):
        ctx=bridge('read_document_context',{})['data']
        selected=next(t for t in ctx['templates'] if t.get('kind')=='related_transaction')
        source=next(s for s in ctx['sources'] if s['id'].startswith('user:'))
        return draft('董事会决议公告',kind='announcement',template_id=selected['id'],
            text='# 董事会决议公告\n一、利润分配\n拟分配2025年净利润，派发现金红利【待补：分红金额】。\n二、战略投资\n拟战略投资100万元。\n【待补：董事会届次和表决结果】',
            basis=[{'source_id':source['id'],'quote':'对国内一家公司战略投资100万','statement':'拟战略投资100万元。'}])
    def notice(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'announcement' if text_only else 'document','reason':'用户要求起草'})
        result=bridge('assess_document_readiness',{'documents':[readiness(build(bridge),gaps)],
            'request_quote':packet['prompt'],'decision':'assess'})
        assert result['terminate'] and result['data']['status']=='waiting_user',result
        emit({'type':'done'})
    runtime.runner = reviewed(notice)
    first=settled(c,request(c,s,'请根据已知情况起草，缺失信息具体标注待补：公司计划分配2025年净利润，召开董事会对国内一家公司战略投资100万。'+('' if text_only else '请制作Word。')))
    assert first['run']['status']=='waiting_user',first
    assert listing(c,s)['items']==[]
    notice_id=first['run']['document_preflight']['notice_id']
    def accept(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'announcement' if text_only else 'document','drafting_notice_id':notice_id,'reason':'用户接受已告知缺口'})
        document=build(bridge)
        opened=bridge('assess_document_readiness',{'documents':[readiness(document,gaps)],'request_quote':packet['prompt'],
            'decision':'proceed_with_gaps','notice_id':notice_id,'choice_quote':packet['prompt']})
        assert ('save_announcement' if text_only else 'make_word') in [t['name'] for t in opened['next_context']['tools']],opened
        args={'documents':[document]}
        if text_only:args['assessment']={'disclosure_needed':'yes','disclosure_scope':'利润分配和战略投资','reason':'用户接受已告知缺口后保存待补正文'}
        bridge('save_announcement' if text_only else 'make_word',args)
        emit({'type':'done'})
    runtime.runner = reviewed(accept)
    out=settled(c,request(c,s,'先按现有资料起草，缺项标注待补。'))
    assert out['run']['status']=='completed',out
    row=listing(c,s)['items'][0]
    text=row['text'] if text_only else document_files.inspect(c.get(row['download']).content)['text']
    assert '2025' in text and '100万元' in text
    assert '【待补：分红金额】' in text and '【待补：董事会届次和表决结果】' in text
    # 公告模板另会为缺失表头字段插入系统占位，故只核对已接受缺口在列。
    assert {'分红金额','董事会届次和表决结果'} <= set(row['pending'])
    assert row['checks']['evidence_bindings']=='recorded_unverified'
    assert c.get('/api/events/'+e['id']).json()==before


def test_declared_gap_must_be_marked_in_the_body(client):
    """A gap list alone does not mark the draft; the placeholder belongs inline."""
    c,runtime,_=client;s,_=new_session(c)
    gap=content_gap('会议日期和表决结果','董事会会议日期与表决结果尚未提供')
    body='# 分析材料\n\n一、当前结论\n会议日期为【待补：会议日期和表决结果】。\n\n二、后续安排\n请核对实施状态。'
    def notice(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','reason':'用户要求制作Word'})
        bridge('read_document_context',{})
        result=bridge('assess_document_readiness',{'documents':[readiness(draft(),[gap])],'request_quote':packet['prompt'],'decision':'assess'})
        assert result['terminate'] and result['data']['status']=='waiting_user',result
        emit({'type':'done'})
    runtime.runner = reviewed(notice)
    first=settled(c,request(c,s))
    assert first['run']['status']=='waiting_user',first
    notice_id=first['run']['document_preflight']['notice_id']
    held=[]
    def accept(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','drafting_notice_id':notice_id,'reason':'用户接受已告知缺口'})
        bridge('read_document_context',{})
        bridge('assess_document_readiness',{'documents':[readiness(draft(),[gap])],'request_quote':packet['prompt'],
                                           'decision':'proceed_with_gaps','notice_id':notice_id,'choice_quote':packet['prompt']})
        missing=bridge('make_word',{'documents':[draft(pending=['会议日期和表决结果'])]})
        held.append(missing['data']['message'])
        result=bridge('make_word',{'documents':[draft(pending=['会议日期和表决结果'],text=body)]})
        assert result['data']['business_state_changed'] is False
        emit({'type':'done'})
    runtime.runner = reviewed(accept)
    out=settled(c,request(c,s,'先按现有资料制作Word，缺项标注待补。'))
    assert out['run']['status']=='completed',out
    assert '正文' in held[0] and '会议日期和表决结果' in held[0]
    row=listing(c,s)['items'][0]
    assert '【待补：会议日期和表决结果】' in document_files.inspect(c.get(row['download']).content)['text']
    assert row['pending']==['会议日期和表决结果']


def test_anchored_revision_preserves_original_without_forced_evidence_review(client):
    c,runtime,_=client;s,_=new_session(c)
    doc=Document();doc.add_paragraph('人工保留段落').runs[0].bold=True
    doc.add_paragraph('金额为1000万元。')
    stream=io.BytesIO();doc.save(stream);raw=stream.getvalue()
    source=legacy_manual_source(runtime,s['id'],raw)
    runtime.semantic_review=lambda rid,run,items,stop: ([{'item_id':r['item_id'],'verdict':'insufficient','reason':'变更金额没有依据'} for r in items],'fixture')
    def runner(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','document_action':'revise','reason':'局部修改'})
        read=bridge('read_document',{'document_id':source['document_id']})['data']
        anchor=next(b for b in read['blocks'] if b['text']=='金额为1000万元。')
        document=draft(title='人工修改稿',document_id=source['document_id'],base_version=1,template_id='source:current',text='',
                       edits=[{'block_id':anchor['id'],'original':'1000','replacement':'987654'}])
        bridge('assess_document_readiness',{'documents':[readiness(document)],'request_quote':packet['prompt'],'decision':'assess'})
        result=bridge('make_word',{'documents':[document]})
        assert result['data']['documents']
        emit({'type':'done'})
    runtime.runner = reviewed(runner)
    out=settled(c,request(c,s,'修订当前Word，缺项标注待补。'))
    assert out['run']['status']=='completed',out
    current=listing(c,s)['items'][0]
    assert current['version']==2 and current['checks']['fact_source_coverage']=='not_run'
    assert '987654' in document_files.inspect(c.get(current['download']).content)['text']
    assert c.get(source['download']).content==raw
    assert not any(e['kind'] in ('document_gap_notice','document_evidence_repair') for e in out['events'])

def test_foreign_document_citation_still_rejected(client):
    c,runtime,_=client;s,_=new_session(c)
    provider(runtime,lambda *_:[draft()]);settled(c,request(c,s));row=listing(c,s)['items'][0]
    other=runtime.store.create_session('chinext','','其他会话',str(uuid4()))
    provider(runtime,lambda *_:[draft(text='以当前资料为准。',basis=[{'statement':'以当前资料为准。',
        'source_id':'document:'+row['document_id']+':1','quote':'以当前资料为准。'}])])
    out=settled(c,request(c,other))
    assert out['run']['status']=='failed'
    assert listing(c,other)['items']==[]


def test_evidence_coverage_preserves_complete_times_tables_and_markers():
    from backend.grounded_text import coverage_items
    text='到账时点为16:00。\n| 投资 | 987654万元 |\n【待补：2025年数据】'
    items=coverage_items(text,[])
    assert [r['value'] for r in items]==text.splitlines()
    assert all(not json.loads(r['context'])['reviewed_statements'] for r in items)

def test_readiness_switch_records_the_stamp_the_worker_reports(client):
    """Production compares phase_started with the run record; an unrecorded switch fails."""
    from fastapi import HTTPException
    c,runtime,_=client;s,_=new_session(c)
    stamps=[]
    def run(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','reason':'用户要求制作Word'})
        opened=bridge('assess_document_readiness',{'documents':[readiness(draft())],'request_quote':packet['prompt'],'decision':'assess'})
        stamps.append(opened['next_context']['system_sha256'])
        bridge('make_word',{'documents':[draft()]})
        emit({'type':'done'})
    runtime.runner = reviewed(run)
    out=settled(c,request(c,s))
    assert out['run']['status']=='completed',out
    assert stamps and runtime.store.run(out['run']['id'])['system_sha256']==stamps[0]
    previous=runtime.runner;runtime.runner=None
    try:
        model=out['run']['model']
        with pytest.raises(HTTPException) as error:
            runtime.receive(out['run']['id'],{'type':'started','system_sha256':'stale-context','model':model['id'],
                                              'provider':model['provider'],'reasoning_effort':model.get('reasoning_effort','off')})
        assert error.value.status_code==409
    finally:
        runtime.runner=previous


def test_saved_draft_keeps_consent_and_identity_for_later_turns(client):
    """The first save binds the plan to the assigned document id without re-asking."""
    c,runtime,_=client;s,_=new_session(c)
    gap=content_gap('会议日期','董事会会议日期尚未提供')
    body='# 分析材料\n\n一、当前结论\n会议日期为【待补：会议日期】。\n\n二、后续安排\n请核对实施状态。'
    def notice(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','reason':'用户要求制作Word'})
        bridge('read_document_context',{})
        result=bridge('assess_document_readiness',{'documents':[readiness(draft(text=body),[gap])],'request_quote':packet['prompt'],'decision':'assess'})
        assert result['terminate'] and result['data']['status']=='waiting_user',result
        emit({'type':'done'})
    runtime.runner = reviewed(notice)
    first=settled(c,request(c,s))
    assert first['run']['status']=='waiting_user',first
    notice_id=first['run']['document_preflight']['notice_id']
    def accept(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','drafting_notice_id':notice_id,'reason':'用户接受已告知缺口'})
        bridge('read_document_context',{})
        opened=bridge('assess_document_readiness',{'documents':[readiness(draft(text=body),[gap])],'request_quote':packet['prompt'],
            'decision':'proceed_with_gaps','notice_id':notice_id,'choice_quote':packet['prompt']})
        assert 'make_word' in [t['name'] for t in opened['next_context']['tools']],opened
        bridge('make_word',{'documents':[draft(text=body)]})
        emit({'type':'done'})
    runtime.runner = reviewed(accept)
    second=settled(c,request(c,s,'先按现有资料制作Word，缺项标注待补。'))
    assert second['run']['status']=='completed',second
    row=listing(c,s)['items'][0]
    stored=runtime.store.run(second['run']['id'])['document_preflight']
    assert stored['documents'][0]['document_id']==row['document_id']
    assert [record['key'] for record in stored['accepted_content']]==[row['document_id']]
    def revise(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','reason':'继续修订同一文稿'})
        bridge('read_document_context',{})
        document=draft(document_id=row['document_id'],base_version=1,text=body+'修订。')
        opened=bridge('assess_document_readiness',{'documents':[readiness(document,[gap])],'request_quote':packet['prompt'],'decision':'assess'})
        assert 'make_word' in [t['name'] for t in opened['next_context']['tools']],opened
        bridge('make_word',{'documents':[document]})
        emit({'type':'done'})
    runtime.runner = reviewed(revise)
    third=settled(c,request(c,s,'请核对当前Word并生成新版本'))
    assert third['run']['status']=='completed',third
    assert listing(c,s)['items'][0]['version']==2


def test_stale_parent_and_manual_full_rewrite_are_rejected(client):
    c,runtime,_=client;s,_=new_session(c)
    provider(runtime,lambda *_:[draft()]);settled(c,request(c,s));row=listing(c,s)['items'][0]
    raw=c.get(row['download']).content
    legacy_manual_source(runtime,s['id'],raw,document_id=row['document_id'],base_version=1)
    provider(runtime,lambda *_:[draft(document_id=row['document_id'],base_version=1)])
    assert settled(c,request(c,s))['run']['status']=='failed'
    provider(runtime,lambda *_:[draft(document_id=row['document_id'],base_version=2)])
    assert settled(c,request(c,s))['run']['status']=='failed'
    assert listing(c,s)['items'][0]['version']==2
    assert c.get(listing(c,s)['items'][0]['download']).content==raw


def test_restart_reconciles_run_but_preserves_saved_document(client):
    from backend.pi_runtime import PiRuntime
    c,runtime,_=client;s,_=new_session(c)
    provider(runtime,lambda *_:[draft()]);result=settled(c,request(c,s));before=listing(c,s)
    runtime.store.update(result['run']['id'],status='running',outcome=None)
    runtime.close()
    recovered=PiRuntime(runtime.root,runtime.directory,runtime.route,config=runtime.config_override,runner=lambda *_:None)
    try:
        recovered.own()
        assert recovered.store.run(result['run']['id'])['status']=='interrupted'
        assert document_store.listing(recovered,s['id'])==before
        path,_=document_store.file(recovered,s['id'],before['items'][0]['document_id'])
        assert path.is_file()
    finally:recovered.close()


def test_document_can_research_but_cannot_mutate_the_knowledge_library(client,monkeypatch):
    from backend import public_sources
    c,runtime,_=client;s,_=new_session(c)
    monkeypatch.setattr(public_sources,'search',lambda *a:{'items':[]})
    def build(ctx,bridge,stop):
        assert bridge('knowledge_web_search',{'query':'公开法规核查'})['data']=={'items':[]}
        from fastapi import HTTPException
        with pytest.raises(HTTPException) as error:bridge('knowledge_propose',{'collection':'laws','items':[]})
        assert error.value.status_code==403
        return [draft()]
    provider(runtime,build)
    assert settled(c,request(c,s))['run']['status']=='completed'


def test_word_return_endpoints_are_retired_without_changing_saved_output(client):
    c,runtime,_=client;s,_=new_session(c)
    provider(runtime,lambda *_:[draft()]);settled(c,request(c,s));before=listing(c,s)
    row=before['items'][0];raw=c.get(row['download']).content
    assert c.post(f'/api/chat/sessions/{s["id"]}/documents/import',json={}).status_code==410
    assert c.post(f'/api/chat/sessions/{s["id"]}/documents/{row["document_id"]}/open',json={}).status_code==410
    assert listing(c,s)==before and c.get(row['download']).content==raw
    assert 'working_file' not in row
    assert not (document_store.folder(runtime,s['id'])/'working').exists()


@pytest.mark.parametrize('external_state',['changed','invalid','missing'])
def test_legacy_working_copy_cannot_change_session_versions(client,external_state):
    c,runtime,_=client;s,_=new_session(c)
    provider(runtime,lambda *_:[draft()]);out=settled(c,request(c,s));row=listing(c,s)['items'][0]
    original=c.get(row['download']).content
    base=document_store.folder(runtime,s['id']);working=base/'working'/'legacy-v1.docx'
    working.parent.mkdir()
    if external_state=='changed':
        doc=Document(io.BytesIO(original));doc.add_paragraph('系统外修改的内容。');doc.save(working)
    elif external_state=='invalid':working.write_bytes(b'saving')
    index=document_store.load(runtime,s['id'])
    index['documents'][0]['versions'][0]['working_file']='working/legacy-v1.docx'
    (base/'index.json').write_text(json.dumps(index,ensure_ascii=False))
    before=(base/'index.json').read_bytes()
    external_bytes=working.read_bytes() if working.exists() else None

    assert listing(c,s)['items'][0]['version']==1
    source=document_runtime.read_document(runtime,out['run'],{'document_id':row['document_id']})
    assert all('系统外修改' not in block['text'] for block in source['blocks'])
    assert (base/'index.json').read_bytes()==before
    assert c.post(f'/api/chat/sessions/{s["id"]}/documents/{row["document_id"]}/review',json={
        'version':1,'sha256':row['sha256'],'decision':'accepted','content_reviewed':True,'visual_reviewed':True}).status_code==200
    def revise(context,bridge,stop):
        read=bridge('read_document',{'document_id':row['document_id']})['data']
        anchor=next(b for b in read['blocks'] if b['text']=='以当前资料为准。')
        return [draft(document_id=row['document_id'],base_version=1,template_id='source:current',text='',
                      edits=[{'block_id':anchor['id'],'original':'以当前资料为准。','replacement':'会话内修订后的内容。'}])]
    provider(runtime,revise)
    assert settled(c,request(c,s,'修改当前文稿并制作Word'))['run']['status']=='completed'
    current=listing(c,s)['items'][0]
    assert current['version']==2 and current['source_type']=='anchored_revision' and 'working_file' not in current
    assert c.get(row['download']).content==original
    assert (working.read_bytes() if working.exists() else None)==external_bytes


def test_changed_registered_file_is_unavailable_and_never_adopted(client):
    c,runtime,_=client;s,_=new_session(c)
    provider(runtime,lambda *_:[draft()]);settled(c,request(c,s));row=listing(c,s)['items'][0]
    base=document_store.folder(runtime,s['id']);before=(base/'index.json').read_bytes()
    path=base/(row['sha256']+'.docx')
    doc=Document(path);doc.add_paragraph('直接改写登记文件。');doc.save(path)
    current=listing(c,s)['items'][0]
    assert current['version']==1 and current['available'] is False
    assert c.get(row['download']).status_code==409
    assert (base/'index.json').read_bytes()==before


def test_financial_formatting_is_preserved_without_claiming_automatic_review(client):
    c,runtime,_=client;s,_=new_session(c)
    text='拟资助1,000.00万元，利率3.00%。'
    runtime.semantic_review=lambda rid,run,items,stop: ([{'item_id':r['item_id'],'verdict':'supported' if r['value']==text else 'insufficient','reason':'隔离用例：数值及单位一致'} for r in items],'fixture')
    provider(runtime,lambda ctx,*_:[draft(text=text,basis=[{'statement':text,'source_id':next(s['id'] for s in ctx['sources'] if s['id'].startswith('user:')),'quote':'拟资助1000万元，利率3%'}])])
    out=settled(c,request(c,s,'拟资助1000万元，利率3%，请制作Word'))
    assert out['run']['status']=='completed',out
    row=listing(c,s)['items'][0]
    assert text in document_files.inspect(c.get(row['download']).content)['text']
    assert row['checks']['fact_source_coverage']=='not_run'


@pytest.mark.parametrize('kind,text_only,declare_kind',[
    ('analysis',False,True), ('announcement',False,True),
    ('announcement',False,False), ('announcement',True,True),
])
def test_complete_draft_is_saved_without_forced_evidence_review(client,kind,text_only,declare_kind):
    c,runtime,_=client;s,_=new_session(c)
    text='\n'.join(f'第{i}项：这是对会话问题的整理与综合分析，待核实的内容保留条件表述。' for i in range(100))[:2422]
    def forbidden_review(*_):
        pytest.fail('制文不得调用自动逐句依据复核')
    runtime.semantic_review=forbidden_review
    skill='name: disclosure-announcement-drafting'
    def runner(packet,emit,bridge,stop):
        args={'domain':'disclosure','intent':'announcement' if text_only else 'document','reason':'完整整理文稿'}
        if declare_kind:args['document_kind']=kind
        routed=bridge('route_request',args)
        assert (skill in routed['next_context']['system']) == (kind=='announcement' and declare_kind)
        ctx=bridge('read_document_context',{})['data']
        tid=next(t['id'] for t in ctx['templates'] if t.get('kind')=='related_transaction') if kind=='announcement' else 'builtin:analysis'
        doc=draft('完整文稿',kind=kind,template_id=tid,text=text)
        doc.pop('basis')  # Non-announcements and announcements may omit bindings.
        opened=bridge('assess_document_readiness',{'documents':[readiness(doc)],'request_quote':packet['prompt'],'decision':'assess'})
        assert (skill in opened['next_context']['system']) == (kind=='announcement')
        production={'documents':[doc]}
        if text_only:production['assessment']={'disclosure_needed':'uncertain','disclosure_scope':'测试事项','reason':'供用户审阅'}
        assert bridge('save_announcement' if text_only else 'make_word',production)['data']['documents']
        emit({'type':'done'})
    runtime.runner=runner
    out=settled(c,request(c,s,'请完整整理为'+('公告正文' if text_only else 'Word')+'。'))
    assert out['run']['status']=='completed',out
    row=listing(c,s)['items'][0]
    snapshot=document_store.snapshot(runtime,s['id'],row['document_id'])
    assert snapshot['text']==text and len(snapshot['text'])==2422
    delivered=row['text'] if text_only else document_files.inspect(c.get(row['download']).content)['text']
    assert text in delivered
    assert row['review_status']=='pending' and row['checks']['fact_source_coverage']=='not_run'
    assert snapshot['semantic_review']['source']=='not_run'
    assert out['run']['document_attempts']==1
    assert not any(e['kind'] in ('document_evidence_repair','text_evidence_review') for e in out['events'])


def test_announcement_revision_loads_skill_from_target_document(client):
    c,runtime,_=client;s,_=new_session(c)
    provider(runtime,lambda ctx,*_:[draft('修订公告',kind='announcement',text='请核对事项实施安排。',
        template_id=next(t['id'] for t in ctx['templates'] if t.get('kind')=='related_transaction'))])
    assert settled(c,request(c,s))['run']['status']=='completed'
    row=listing(c,s)['items'][0]
    def runner(packet,emit,bridge,stop):
        routed=bridge('route_request',{'domain':'disclosure','intent':'document','document_action':'revise',
            'target_document_id':row['document_id'],'reason':'修改已有公告Word'})
        assert 'name: disclosure-announcement-drafting' in routed['next_context']['system']
        doc=draft(row['title'],kind='announcement',document_id=row['document_id'],base_version=1,
                  template_id='source:current',text='',pending=row['pending'])
        read=bridge('read_document',{'document_id':row['document_id']})['data']
        block=next(b for b in read['blocks'] if b['text']=='请核对事项实施安排。')
        doc['edits']=[{'block_id':block['id'],'original':block['text'],'replacement':'请核对最新实施安排。'}]
        bridge('assess_document_readiness',{'documents':[readiness(doc,[content_gap(p) for p in row['pending']])],
            'request_quote':packet['prompt'],'decision':'draft_with_placeholders','choice_quote':packet['prompt']})
        bridge('make_word',{'documents':[doc]})
        emit({'type':'done'})
    runtime.runner=runner
    out=settled(c,request(c,s,'将公告中“请核对事项实施安排。”改为“请核对最新实施安排。”，先制作Word，缺项标注待补。'))
    assert out['run']['status']=='completed',out
    current=listing(c,s)['items'][0]
    assert current['version']==2 and current['checks']['source_package_preserved']
    assert '请核对最新实施安排。' in document_files.inspect(c.get(current['download']).content)['text']
