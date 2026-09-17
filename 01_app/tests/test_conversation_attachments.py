"""Attachment parsing/storage/scope checks; no model or business dialogue is run."""
import base64
import io
from uuid import uuid4
import pytest
from docx import Document
from openpyxl import Workbook
from fastapi import HTTPException
from backend import conversation_attachments as attachments, document_store, document_runtime
from test_pi_runtime import client


def encoded(document):
    stream=io.BytesIO();document.save(stream);return stream.getvalue()


def session(runtime):
    return runtime.store.create_session('chinext',None,'附件读取检查',str(uuid4()))


def upload(c,s,raw,name):
    return c.post(f'/api/chat/sessions/{s["id"]}/attachments',json={'filename':name,'content_base64':base64.b64encode(raw).decode()})


def reader_run(runtime,s,ids):
    run,_=runtime.store.accept(s,{'text':'读取补充资料','stage':'auto','request_id':str(uuid4()),'attachment_ids':ids},{})
    return runtime.store.update(run['id'],stage='document',intent='document',status='running')


def test_docx_original_and_locators_survive_repeated_upload(client):
    c,runtime,_=client;s=session(runtime)
    doc=Document();doc.add_paragraph('新任董事会秘书：张示例。')
    doc.add_table(rows=1,cols=2).rows[0].cells[0].text='任职日期'
    doc.tables[0].rows[0].cells[1].text='2026年9月'
    raw=encoded(doc);response=upload(c,s,raw,'补充说明.docx')
    assert response.status_code==200,response.text
    row=response.json();assert upload(c,s,raw,'补充说明.docx').json()['id']==row['id']
    run=reader_run(runtime,s,[row['id']]);first=attachments.read(runtime,run['id'],{'attachment_id':row['id'],'limit':1})['data']
    assert first['next_offset']==1
    assert '张示例' in first['blocks'][0]['text']
    source=first['blocks'][0]['source_id']
    assert attachments.source_text(runtime,runtime.store.run(run['id']),source)==first['blocks'][0]['text']
    fresh=runtime.store.run(run['id']);context=document_runtime.context(runtime,fresh)
    assert context['attachments'][0]['sha256']==row['sha256']
    assert document_runtime.source_text(runtime,fresh,context,source)==first['blocks'][0]['text']
    assert 'read_attachment' in [t['name'] for t in runtime.tools('document',context)]
    assert (attachments.directory(runtime,s['id'],row['id'])/'original.docx').read_bytes()==raw
    assert document_store.listing(runtime,s['id'])['items']==[]


def test_excel_preserves_raw_values_formula_warning_and_row_address(client):
    c,runtime,_=client;s=session(runtime)
    book=Workbook();sheet=book.active;sheet.title='补充数据'
    sheet.append(['姓名','投资金额（万元）']);sheet.append(['示例',12.34567]);sheet['B2'].number_format='0.00'
    sheet.append(['合计','=B2'])
    row=upload(c,s,encoded(book),'补充数据.xlsx').json();run=reader_run(runtime,s,[row['id']])
    data=attachments.read(runtime,run['id'],{'attachment_id':row['id']})['data']
    second=data['blocks'][1]
    assert second['locator']=='sheet:补充数据/row:2'
    assert second['cell_metadata'][1]['coordinate']=='B2'
    assert second['cell_metadata'][1]['raw_value']=='12.34567'
    assert '12.34567' in second['text']
    assert data['blocks'][2]['cell_metadata'][1]['formula']=='B2'
    assert any('缓存' in w for w in data['attachment']['warnings'])


def test_unselected_cross_session_and_tampered_files_cannot_be_read(client):
    c,runtime,_=client;s=session(runtime);other=session(runtime)
    doc=Document();doc.add_paragraph('仅当前会话资料')
    row=upload(c,s,encoded(doc),'资料.docx').json()
    run=reader_run(runtime,s,[])
    with pytest.raises(HTTPException) as exc:attachments.read(runtime,run['id'],{'attachment_id':row['id']})
    assert exc.value.status_code==403
    with pytest.raises(HTTPException) as exc:attachments.validate_selection(runtime,other['id'],[row['id']])
    assert exc.value.status_code==404
    runtime.store.update(run['id'],attachment_ids=[row['id']])
    (attachments.directory(runtime,s['id'],row['id'])/'original.docx').write_bytes(b'changed isolated file')
    with pytest.raises(HTTPException) as exc:attachments.read(runtime,run['id'],{'attachment_id':row['id']})
    assert exc.value.status_code==409


def test_format_path_and_archive_guards(client):
    c,runtime,_=client;s=session(runtime)
    doc=Document();doc.add_paragraph('普通材料');raw=encoded(doc)
    assert upload(c,s,raw,'../资料.docx').status_code==422
    assert upload(c,s,raw,'错误类型.xlsx').status_code==422
    assert upload(c,s,b'not a zip','损坏.docx').status_code==422
    runtime.store.edit_session(s['id'],archived=True)
    assert upload(c,s,raw,'归档后.docx').status_code==409
