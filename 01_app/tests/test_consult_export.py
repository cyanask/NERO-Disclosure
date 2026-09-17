"""Consultation work-draft export and layered assessment facts."""
import hashlib
import io
from pathlib import Path
import pytest
from docx import Document
from fastapi import HTTPException
from backend import consult_export, disclosure_contract as contract

ROOT=Path(__file__).resolve().parents[1]
LAYOUT=consult_export.layout(ROOT,'chinext')


def test_render_keeps_displayed_content_and_never_looks_like_a_formal_notice():
    raw=consult_export.render('# 一、结论\n董事会决议涉及应披露事项时须公告。\n\n- 第一点\n- 第二点\n\n| 项目 | 金额 |\n| --- | --- |\n| 收入 | 100 |\n',
        '测试标题','测试公司',LAYOUT,'2026-09-16 10:00')
    doc=Document(io.BytesIO(raw))
    texts=[p.text for p in doc.paragraphs]
    assert '咨询工作稿 · 待复核' in texts
    assert '测试标题' in texts
    assert '董事会决议涉及应披露事项时须公告。' in texts
    assert any(text.startswith('· 第一点') for text in texts)
    assert any('咨询工作稿／待复核' in text for text in texts)
    assert doc.tables and doc.tables[0].cell(1,1).text=='100'
    sizes={round(run.font.size.pt,1) for p in doc.paragraphs for run in p.runs if run.text.strip()}
    assert sizes <= {12.0,14.0,10.5,9.0}, sizes


class Store:
    def __init__(self,run,journal,sid):self._run=run;self._journal=journal;self._sid=sid
    def run(self,rid):
        if rid!=self._run['id']:raise HTTPException(404,'执行不存在')
        return self._run
    def journal(self,rid):return self._journal
    def session(self,sid):
        if sid!=self._sid:raise HTTPException(404,'会话不存在')
        return {'id':sid,'board':'chinext','event_id':'conversation:'+sid,'title':'测试咨询','archived':0}


def runtime(tmp_path,run,journal,sid='sid-1'):
    class Value:
        root=tmp_path
        store=Store(run,journal,sid)
    return Value()


def settled_run(sid='sid-1'):return {'id':'run-1','session_id':sid,'status':'completed','stage':'chat'}


def test_retired_export_refuses_creation_but_keeps_historical_files(tmp_path):
    import json
    from uuid import uuid4
    value=runtime(tmp_path,settled_run(),[])
    with pytest.raises(HTTPException) as error:
        consult_export.create(value,'sid-1',{'run_id':'run-1','text':'旧回复'})
    assert error.value.status_code==410
    raw=consult_export.render('历史文稿','历史文件','测试公司',LAYOUT,'2026-09-16')
    sha=hashlib.sha256(raw).hexdigest();identity=str(uuid4())
    folder=tmp_path/'work/consult-exports/sid-1';folder.mkdir(parents=True)
    (folder/(sha+'.docx')).write_bytes(raw)
    (folder/(identity+'.json')).write_text(json.dumps({'id':identity,'session_id':'sid-1','sha256':sha,'filename':'历史文件.docx'}))
    assert consult_export.listing(value,'sid-1')['items'][0]['id']==identity
    path,filename=consult_export.file_path(value,'sid-1',identity)
    assert path.read_bytes()==raw and filename=='历史文件.docx'


def test_export_refuses_text_that_was_never_displayed(tmp_path):
    journal=[{'kind':'assistant','body':{'phase':'final','text':'已显示的答复'}}]
    value=runtime(tmp_path,settled_run(),journal)
    with pytest.raises(HTTPException) as error:
        consult_export.create(value,'sid-1',{'run_id':'run-1','text':'使用者自己改写的另一段文字'})
    assert error.value.status_code==410


def test_export_refuses_a_still_running_round(tmp_path):
    run={**settled_run(),'status':'running'}
    value=runtime(tmp_path,run,[{'kind':'assistant','body':{'text':'部分内容'}}])
    with pytest.raises(HTTPException) as error:
        consult_export.create(value,'sid-1',{'run_id':'run-1','text':'部分内容'})
    assert error.value.status_code==410


def test_title_prefers_heading_then_session_title():
    assert consult_export.title_of('# 一、结论\n可以。\n后续说明', '会话标题')=='一、结论'
    assert consult_export.title_of('可以。当前绑定范围为创业板。', '会话标题')=='会话标题'
    assert consult_export.title_of('可以。当前绑定范围为创业板。', '新会话')=='可以。当前绑定范围为创业板。'


def test_export_uses_the_leading_heading_as_its_single_document_title():
    text='# 关联方财务资助事项合规分析工作稿\n\n## 一、结论\n事实仍待核实。\n\n## 二、依据\n须对照相关规则判断。'
    title=consult_export.title_of(text,'原会话标题')
    raw=consult_export.render(text,title,'测试公司',LAYOUT,'2026-09-16 18:00')
    doc=Document(io.BytesIO(raw))
    texts=[p.text for p in doc.paragraphs]
    assert texts.count(title)==1
    assert [p.style.name for p in doc.paragraphs if p.text==title]==['Title']
    assert texts[-5:]==[title,'一、结论','事实仍待核实。','二、依据','须对照相关规则判断。']


LAW={'id':'law-1','title':'测试规则','article':'第一条','source_kind':'official_rule','text':'第一条 发生重大事项应当披露。'}


def unreadable(identity):
    raise ValueError('该原件在本测试中不可读取')


def assessment(summary,facts,paths,history_items=None):
    event={'id':'e1','company_id':'c1','layer':'chinext','company_name':'测试公司','stock_code':'300001',
           'intake_mode':'open','facts':{'event_date':'2026-01-01'},'summary':summary}
    catalog={'sources':[LAW],'client_history':{'items':history_items or []}}
    result={'source_ids':['law-1'],'facts':facts,'matters':[{'matter_id':'m1','duty_status':'mandatory',
        'reasoning_items':paths,'decisive_questions':[],'reassessment_conditions':['条件变化时重判'],
        'timing_status':'pending','trigger_events':[],'deadline_basis':[],'historical_links':[],
        'history_status':'not_connected','comparable_cases':[],'urgency':'normal','specialist_required':False,
        'special_review':'not_required','calculations':[]}]}
    errors,warnings,review=contract.assessment_checks(event,catalog,result,history_reader=unreadable)
    return {i['code'] for i in errors},{i['code'] for i in warnings}


MISSING=[{'id':'announcement-missing','sha256':'x'}]


def test_non_decisive_history_gap_is_a_pending_note_not_a_blocker():
    paths=[{'source_id':'law-1','locator':'第一条','quote':'发生重大事项应当披露','fact_keys':['key_a'],'outcome':'established'}]
    facts=[{'key':'key_a','value':'发生重大事项','status':'user_statement','source_ref':'summary','quote':'公司发生重大事项'},
           {'key':'h1','value':'2022年曾召开董事会','status':'user_statement','source_ref':'announcement-missing','quote':''}]
    errors,warnings=assessment('公司发生重大事项，应当披露。',facts,paths,MISSING)
    assert 'history_fact_unlocatable' not in errors
    assert 'fact_pending_supplement' in warnings


def test_decisive_history_fact_still_blocks_the_node():
    paths=[{'source_id':'law-1','locator':'第一条','quote':'发生重大事项应当披露','fact_keys':['h1'],'outcome':'established'}]
    facts=[{'key':'h1','value':'2022年曾召开董事会','status':'user_statement','source_ref':'announcement-missing','quote':''}]
    errors,_=assessment('公司事项情况见材料。',facts,paths,MISSING)
    assert 'history_fact_unlocatable' in errors


def test_summary_fact_binds_to_the_users_own_sentence():
    paths=[{'source_id':'law-1','locator':'第一条','quote':'发生重大事项应当披露','fact_keys':['key_a'],'outcome':'established'}]
    facts=[{'key':'key_a','value':'重大事项','status':'user_statement','source_ref':'summary','quote':''}]
    errors,warnings=assessment('公司发生重大事项，应当披露。',facts,paths)
    assert 'fact_source_mismatch' not in errors
    assert 'fact_quote_bound' in warnings
    assert facts[0]['quote']=='公司发生重大事项，应当披露。'
