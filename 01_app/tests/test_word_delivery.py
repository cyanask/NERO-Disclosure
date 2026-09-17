from internal_workflow import context as read_context
"""External Word round trips in isolated public fixtures; no Office/visual acceptance."""
import base64
import hashlib
import io
from docx import Document
from docx.shared import Pt
from test_agent_tasks import env,event,claimed,post,advance,submit,latest,confirm_stage,plan_candidate,approve_template
from scripts.word_renderer import render


def ready(c,token):
    from conftest import ready_layout
    ready_layout(c.app.state.pi_runtime.root)
    e=event(c);task,lease=claimed(c,e,token)
    assert submit(c,e,task,lease,token).status_code==200
    assert post(c,e,f"agent-tasks/{task['id']}/adopt",token).status_code==200
    assert advance(c,e,'assessment').status_code==200
    assert confirm_stage(c,e,'assessment','prepare_mandatory').status_code==200
    task,lease=claimed(c,e,token,'plan')
    ctx=read_context(c,e,task,token).json()
    assert ctx['verification_gate']['status']=='PASS',ctx['verification_gate']
    plan=plan_candidate(ctx);items=plan['items']
    assert submit(c,e,task,lease,token,plan).status_code==200
    assert post(c,e,f"agent-tasks/{task['id']}/adopt",token).status_code==200
    assert advance(c,e,'plan').status_code==200
    assert confirm_stage(c,e,'plan').status_code==200
    approve_template(c,e,token,ctx,plan)
    task,lease=claimed(c,e,token,'draft')
    body='证券代码：000000 证券简称：模拟公司 主办券商：模拟券商 公告编号：模拟001\n# 董事会决议公告\n'
    body+='\n'.join('## '+i['title']+'\n模拟金额 100.00 元。' for i in items)
    body+='\n| 项目 | 金额（元） |\n| --- | --- |\n| 本次 | 100.00 |'
    assert submit(c,e,task,lease,token,{'template_id':ctx['templates'][0]['id'],'text':body,'requirement_map':{r['requirement_id']:'模拟金额 100.00 元。' for r in plan['requirements']}}).status_code==200
    r=post(c,e,f"agent-tasks/{task['id']}/adopt",token);assert r.status_code==200,r.text
    assert advance(c,e,'draft').status_code==200
    packet=c.get('/api/events/'+e['id']+'/word/context').json()
    return e,packet


def word_bytes(root,packet):
    template=root/'external-template.docx';template.write_bytes(base64.b64decode(packet['template_base64']))
    raw,_=render(packet['event'],packet['layout'],template)
    return raw


def register(c,e,packet,raw,**extra):
    return post(c,e,'artifacts',input_fingerprint=packet['input_fingerprint'],sha256=hashlib.sha256(raw).hexdigest(),filename='模拟验收.docx',content_base64=base64.b64encode(raw).decode(),host_qa='自动测试资料，无人工或视觉验收',**extra)


def test_external_word_roundtrip_and_immutable_history(env):
    c,tokens,root,_=env;e,packet=ready(c,tokens[0]);raw=word_bytes(root,packet)
    r=register(c,e,packet,raw);assert r.status_code==200,r.text
    artifact=r.json()['artifacts'][0];assert artifact['verification']['visual_acceptance']=='not_verified'
    assert r.json()['stage']=='awaiting_word_confirmation'
    assert len(register(c,e,packet,raw).json()['artifacts'])==1
    url=f"/api/events/{e['id']}/artifacts/{artifact['id']}/file"
    assert c.get(url).content==raw
    current=latest(c,e);assert c.patch('/api/events/'+e['id'],json={'expected_revision':current['revision'],'summary':'模拟变更'}).status_code==200
    assert register(c,e,packet,raw).status_code==409
    assert c.get(url).content==raw  # retained original history, not current acceptance


def test_word_modified_number_font_and_hash_rejected(env):
    c,tokens,root,_=env;e,packet=ready(c,tokens[0]);raw=word_bytes(root,packet)
    for change in ('number','font'):
        doc=Document(io.BytesIO(raw))
        if change=='number':doc.tables[0].cell(1,1).text='999.00'
        else:doc.paragraphs[0].runs[0].font.size=Pt(32)
        out=io.BytesIO();doc.save(out)
        assert register(c,e,packet,out.getvalue()).status_code==422
    assert post(c,e,'artifacts',input_fingerprint=packet['input_fingerprint'],sha256='0'*64,filename='file.docx',content_base64=base64.b64encode(raw).decode()).status_code==422
    assert not latest(c,e).get('artifacts')


def test_backend_does_not_import_external_renderer(env):
    c,tokens,root,_=env;e=event(c)
    assert post(c,e,'word/prepare').status_code==410
    assert c.get('/api/events/'+e['id']+'/word/context').status_code==409
    from pathlib import Path
    assert 'word_renderer' not in (Path(__file__).resolve().parents[1]/'backend/app.py').read_text()


def test_external_relationship_and_old_template_are_blocked(env):
    import zipfile
    c,tokens,root,_=env;e,packet=ready(c,tokens[0]);raw=word_bytes(root,packet)
    output=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(output,'w') as dst:
        for name in src.namelist():
            data=src.read(name)
            if name=='word/_rels/document.xml.rels':data=data.replace(b'</Relationships>',b"<Relationship Id='rIdExternalTest' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink' Target='https://example.invalid/' TargetMode = 'External'/></Relationships>")
            dst.writestr(name,data)
    assert register(c,e,packet,output.getvalue()).status_code==422
    import json
    template=next(t for t in json.loads((root/'templates/boards/chinext/manifest.json').read_text()) if t['id']==packet['event']['draft']['template_id'])
    path=root/'templates'/template['file'];path.write_bytes(path.read_bytes()+b'changed')
    assert c.get('/api/events/'+e['id']+'/word/context').status_code==409
    assert register(c,e,packet,raw).status_code==409
