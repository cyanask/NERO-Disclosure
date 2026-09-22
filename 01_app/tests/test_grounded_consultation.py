"""Real runtime and Word subprocess, isolated data and explicit reviewer verdicts."""
import json
import threading
from uuid import uuid4
from urllib.parse import parse_qs,urlsplit
import pytest
from fastapi import HTTPException
from backend import public_sources,document_store,document_files,reply_document
from test_pi_runtime import client,settled
from test_document_runtime import request,draft,readiness,listing


def session(runtime):
    return runtime.store.create_session('chinext','','依据与制文测试',str(uuid4()))


def review_only(runtime,allowed):
    # Deliberate test oracle. It asserts support only for these fixture sentences,
    # never substitutes for the production model review or its parser.
    def review(rid,run,items,stop):
        return [{'item_id':r['item_id'],'verdict':'supported' if r['value'] in allowed else 'insufficient',
                 'reason':'预先给定的隔离用例判定','locator':'fixture'} for r in items],'fixture'
    runtime.semantic_review=review


def test_search_has_no_domain_filter_and_keeps_candidate_status(monkeypatch):
    seen=[]
    def fetch(url,search=False):
        seen.append(url)
        return b'<rss><channel><item><title>Rule</title><link>https://www.chinaclear.cn/test</link><description>candidate</description></item><item><title>Commentary</title><link>https://example.org/article</link></item></channel></rss>','text/xml',url
    monkeypatch.setattr(public_sources,'fetch',fetch)
    result=public_sources.search('中国结算 权益分派','chinext')
    assert parse_qs(urlsplit(seen[0]).query)['q']==['中国结算 权益分派']
    assert len(result['items'])==2
    assert '尚未核对' in result['coverage']
    assert result['channel']=='bing_rss' and len(seen)==2


def test_search_prefers_the_html_channel_and_unwraps_redirects(monkeypatch):
    page="""<tr><td>1.</td><td><a rel="nofollow" href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fwww.chinaclear.cn%2Fzdjs%2Fiszsc.pdf&amp;rut=a" class='result-link'>证券发行人业务指南</a></td></tr>
<tr><td class='result-snippet'>深圳分公司业务指南</td></tr>
<tr><td>2.</td><td><a class="result__a" href="https://www.szse.cn/direct.html">Direct entry</a></td></tr>"""
    monkeypatch.setattr(public_sources,'fetch',lambda url,search=False:(page.encode(),'text/html',url))
    result=public_sources.search('权益分派 业务指南','chinext')
    assert result['channel']=='duckduckgo_html'
    assert [x['url'] for x in result['items']]==['https://www.chinaclear.cn/zdjs/iszsc.pdf','https://www.szse.cn/direct.html']
    assert result['items'][0]['title']=='证券发行人业务指南' and result['items'][0]['snippet']=='深圳分公司业务指南'


def test_search_falls_through_channels_without_inventing_results(monkeypatch):
    def fetch(url,search=False):
        return (b'<html><body>no results</body></html>','text/html',url) if 'duckduckgo' in url else (b'<rss><channel></channel></rss>','text/xml',url)
    monkeypatch.setattr(public_sources,'fetch',fetch)
    result=public_sources.search('权益分派 业务指南','chinext')
    assert result['strategy']=='official_navigation_fallback' and all(x.get('kind') for x in result['items'])


def test_search_retries_a_throttled_primary_channel(monkeypatch):
    calls=[]
    page="""<tr><td>1.</td><td><a href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fwww.neeq.com.cn%2Flaw" class='result-link'>Rule</a></td></tr>"""
    def fetch(url,search=False):
        calls.append(url)
        if len(calls)==1:raise HTTPException(status_code=409,detail='公开来源返回 HTTP 202')
        return page.encode(),'text/html',url
    monkeypatch.setattr(public_sources,'fetch',fetch)
    monkeypatch.setattr(public_sources,'SEARCH_RETRY_PAUSE',0)
    result=public_sources.search('权益分派','chinext')
    assert result['channel']=='duckduckgo_html' and result['items'][0]['url']=='https://www.neeq.com.cn/law'


@pytest.mark.parametrize('url',['https://www.chinaclear.cn/guide','https://example.org/research'])
def test_download_accepts_public_hosts_without_allowlist(monkeypatch,url):
    seen=[]
    class Socket:
        def close(self):pass
    class TLS:
        def wrap_socket(self,s,server_hostname):return s
    class Response:
        status=200
        def read(self,n):return b'<html>original</html>'
        def getheader(self,k,d=''):return 'text/html'
    class Connection:
        sock=None
        def __init__(self,*args,**kw):seen.append(args[0])
        def request(self,*args,**kw):pass
        def getresponse(self):return Response()
        def close(self):pass
    monkeypatch.setattr(public_sources,'public_addresses',lambda host:['1.1.1.2'])
    monkeypatch.setattr(public_sources.ssl,'create_default_context',lambda:TLS())
    monkeypatch.setattr(public_sources.socket,'create_connection',lambda *a,**kw:Socket())
    monkeypatch.setattr(public_sources.http.client,'HTTPSConnection',Connection)
    assert public_sources.fetch(url)[0]==b'<html>original</html>'
    assert seen==[urlsplit(url).hostname]


@pytest.mark.parametrize('address',['127.0.0.1','10.0.0.1','::1','169.254.169.254'])
def test_download_still_blocks_private_addresses(monkeypatch,address):
    monkeypatch.setattr(public_sources.socket,'getaddrinfo',lambda *a,**kw:[(None,None,None,None,(address,443))])
    with pytest.raises(HTTPException) as error:public_sources.fetch('https://arbitrary.example/data')
    assert error.value.status_code==403


def test_ungrounded_answer_is_published_and_completes(client):
    c,runtime,_=client;s=session(runtime)
    def runner(packet,emit,bridge,stop):
        routed=bridge('route_request',{'domain':'disclosure','intent':'consult','reason':'咨询'})
        assert routed['next_context']['requires_result'] is False
        assert 'request_information' in [t['name'] for t in routed['next_context']['tools']]
        emit({'type':'text_delta','message':0,'delta':'无依据的16:00断言'})
        emit({'type':'assistant','text':'无依据的16:00断言','message':0,'stopReason':'stop'})
        emit({'type':'done'})
    runtime.runner=runner
    out=settled(c,request(c,s,'查询办理时限'))
    assert out['run']['status']=='completed',out['run'].get('reason')
    assert not any(e['kind']=='text_delta' for e in out['events'])
    assert [e['body']['text'] for e in out['events'] if e['kind']=='assistant']==['无依据的16:00断言']


def test_consultation_grounded_answer_and_render_exact_reply(client):
    c,runtime,_=client;s=session(runtime)
    text='测试规则要求在16:00前完成申报。'
    review_only(runtime,{text})
    # Use a registered source resolver stub, never a real regulation or live model.
    original=runtime.operation
    runtime.operation=lambda rid,name,args: {'text':text,'source_kind':'official_rule','url':'https://example.org/fixture'} if name=='library.read' else original(rid,name,args)
    def consult(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'consult','reason':'咨询'})
        result=bridge('submit_consultation',{'text':text,'basis':[{'statement':text,'source_id':'library:fixture','quote':text}]})
        assert result['terminate']
        emit({'type':'assistant','message':0,'text':'追加一条未经核实的新结论','stopReason':'stop'})
        emit({'type':'done'})
    runtime.runner=consult
    first=settled(c,request(c,s,'查询办理时限'))
    assert first['run']['status']=='completed',first['run'].get('reason')
    answers=[e for e in first['events'] if e['kind']=='assistant']
    assert [e['body']['text'] for e in answers]==[text]
    assert not any(e['kind']=='document_gap_notice' for e in first['events'])
    source={'reply_run_id':first['run']['id'],'reply_message_seq':answers[0]['seq']}
    def render(packet,emit,bridge,stop):
        routed=bridge('route_request',{'domain':'disclosure','intent':'document','document_kind':'analysis','document_action':'render','reason':'原回复转Word',**source})
        assert {'make_word','read_document','knowledge_web_search'} <= {t['name'] for t in routed['next_context']['tools']}
        assert routed['next_context']['stage']=='document'
        with pytest.raises(HTTPException):bridge('make_word',{'documents':[draft(text='擅自改写')]})
        result=bridge('make_word',{})
        assert result['data']['documents']
        emit({'type':'done'})
    runtime.runner=render
    second=settled(c,request(c,s,'把刚才回复制作成Word'))
    assert second['run']['status']=='completed',second['run'].get('reason')
    row=listing(c,s)['items'][0]
    snapshot=document_store.snapshot(runtime,s['id'],row['document_id'])
    assert snapshot['text']==text and snapshot['source_reply']['reviewed']
    assert text in document_files.inspect(c.get(row['download']).content)['text']
    assert '核实数值' not in json.dumps(snapshot,ensure_ascii=False)
    assert not any(e['kind']=='document_gap_notice' for e in second['events'])


def test_invalid_binding_and_missing_review_only_add_reminders(client):
    c,runtime,_=client;s=session(runtime)
    review_only(runtime,{'不应展示'})
    def runner(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'consult','reason':'咨询'})
        reply=bridge('submit_consultation',{'text':'不应展示','basis':[{'statement':'不应展示','source_id':'user:0','quote':'伪造原句'}]})
        assert reply['data']['status']=='completed' and reply['terminate'] and reply['data']['warnings']
        emit({'type':'done'})
    runtime.runner=runner
    out=settled(c,request(c,s,'咨询'))
    assert out['run']['status']=='completed'
    answers=[e for e in out['events'] if e['kind']=='assistant']
    assert [e['body']['text'] for e in answers]==['不应展示'] and answers[0]['body']['evidence_status']=='unverified'
    assert any(e['kind']=='consultation_evidence_warnings' for e in out['events'])
    # A reviewer that cannot answer is a reminder as well; it never withholds the text.
    text='测试规则要求在16:00前完成申报。'
    original=runtime.operation
    runtime.operation=lambda rid,name,args: {'text':text,'source_kind':'official_rule','url':'https://example.org/fixture'} if name=='library.read' else original(rid,name,args)
    runtime.semantic_review=lambda *a:(None,'invalid_result')
    other=session(runtime)
    def broken(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'consult','reason':'咨询'})
        reply=bridge('submit_consultation',{'text':text,'basis':[{'statement':text,'source_id':'library:fixture','quote':text}]})
        assert reply['terminate'] and any('依据复核未完成' in w['reason'] for w in reply['data']['warnings'])
        emit({'type':'done'})
    runtime.runner=broken
    second=settled(c,request(c,other,'咨询'))
    assert second['run']['status']=='completed'
    assert [e['body']['text'] for e in second['events'] if e['kind']=='assistant']==[text]


def test_evidence_reminders_name_the_exact_failure(client):
    c,runtime,_=client;s=session(runtime)
    text='测试规则要求在16:00前完成申报。'
    original=runtime.operation
    runtime.operation=lambda rid,name,args: {'text':text,'source_kind':'official_rule','url':'https://example.org/fixture'} if name=='library.read' else original(rid,name,args)
    def runner(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'consult','reason':'咨询'})
        reply=bridge('submit_consultation',{'text':text,'basis':[
            {'statement':text,'source_id':'library:fixture','quote':'来源中不存在的原句'},
            {'statement':'正文中不存在的断言','source_id':'library:fixture','quote':text}]})
        reasons=[w['reason'] for w in reply['data']['warnings']]
        assert any('引文在所引来源中定位不到' in r for r in reasons)
        assert any('未逐字出现在提交正文中' in r for r in reasons)
        assert reply['terminate'] and reply['data']['status']=='completed'
        emit({'type':'done'})
    runtime.runner=runner
    out=settled(c,request(c,s,'查询办理时限'))
    assert out['run']['status']=='completed'


def test_word_evidence_does_not_block_save_or_create_numeric_gap(client):
    c,runtime,_=client;s=session(runtime)
    review_only(runtime,set())
    def runner(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','document_action':'create','reason':'研究后起草'})
        doc=draft(text='必须在16:00前支付987654万元。')
        bridge('assess_document_readiness',{'documents':[readiness(doc)],'request_quote':packet['prompt'],'decision':'assess'})
        result=bridge('make_word',{'documents':[doc]})
        assert result['data']['documents']
        emit({'type':'done'})
    runtime.runner=runner
    out=settled(c,request(c,s,'研究后制作Word，缺项标注待补。'))
    assert out['run']['status']=='completed',out
    assert not any(e['kind']=='document_gap_notice' for e in out['events'])
    row=listing(c,s)['items'][0]
    assert document_store.snapshot(runtime,s['id'],row['document_id'])['text']=='必须在16:00前支付987654万元。'
    assert row['checks']['fact_source_coverage']=='not_run'
    assert not any(e['kind']=='document_evidence_repair' for e in out['events'])


def test_reply_binding_rejects_other_session_and_cancelled_render(client):
    c,runtime,_=client;a=session(runtime);b=session(runtime)
    run,_=runtime.store.accept(a,{'text':'原咨询','stage':'chat','request_id':str(uuid4())},{})
    runtime.store.update(run['id'],status='completed')
    runtime.trace(run['id'],'assistant',{'text':'历史回复：16:00。','stopReason':'stop'})
    seq=runtime.store.journal(run['id'])[-1]['seq']
    with pytest.raises(HTTPException):reply_document.bind(runtime,{'id':'new','session_id':b['id']},run['id'],seq)
    result=reply_document.bind(runtime,{'id':'new','session_id':a['id']},run['id'],seq)
    assert not result['reviewed'] and result['text']=='历史回复：16:00。'


def test_word_save_preserves_existing_gap_consent_without_evidence_repair(client):
    c,runtime,_=client;s=session(runtime)
    fact='测试规则要求在16:00前完成申报。';marker='【待补：实施参数】'
    review_only(runtime,{fact,marker})
    original=runtime.operation
    runtime.operation=lambda rid,name,args: {'text':fact,'source_kind':'official_rule'} if name=='library.read' else original(rid,name,args)
    def runner(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','document_action':'create','reason':'制作待补稿'})
        doc=draft(text=marker+'\n必须在15:00前到账。',pending=['实施参数'])
        gap={'label':'实施参数','category':'content','state':'missing','reason':'未提供实施参数'}
        bridge('assess_document_readiness',{'documents':[readiness(doc,[gap])],'request_quote':packet['prompt'],
            'decision':'draft_with_placeholders','choice_quote':packet['prompt']})
        assert bridge('make_word',{'documents':[doc]})['data']['documents']
        emit({'type':'done'})
    runtime.runner=runner
    out=settled(c,request(c,s,'先按现有资料制作Word，缺项标注待补。'))
    assert out['run']['status']=='completed',out['run']
    assert len([e for e in out['events'] if e['kind']=='document_preflight_registered'])==1
    assert not any(e['kind']=='document_gap_notice' for e in out['events'])
    row=listing(c,s)['items'][0]
    assert row['pending']==['实施参数']
    assert marker+'\n必须在15:00前到账。' in document_files.inspect(c.get(row['download']).content)['text']
    assert out['run']['document_attempts']==1
    assert not any(e['kind']=='document_evidence_repair' for e in out['events'])
    assert out['run']['document_preflight']['consent']['quote']=='先按现有资料制作Word，缺项标注待补。'


def test_cancelled_reply_render_does_not_publish(client):
    c,runtime,_=client;s=session(runtime)
    prior,_=runtime.store.accept(s,{'text':'原咨询','stage':'chat','request_id':str(uuid4())},{})
    runtime.store.update(prior['id'],status='completed')
    runtime.trace(prior['id'],'assistant',{'text':'历史回复：16:00。','stopReason':'stop'})
    seq=runtime.store.journal(prior['id'])[-1]['seq']
    run,_=runtime.store.accept(s,{'text':'转Word','stage':'document','request_id':str(uuid4())},{})
    source=reply_document.bind(runtime,run,prior['id'],seq)
    runtime.store.update(run['id'],document_action='render',reply_source=source)
    stop=threading.Event();stop.set()
    with pytest.raises(HTTPException):reply_document.render(runtime,run['id'],{},stop)
    assert listing(c,s)['items']==[]


def test_confirmed_model_generated_analysis_requires_anchors(client):
    import io
    from docx import Document
    from test_document_runtime import legacy_manual_source
    c,runtime,_=client;s=session(runtime)
    doc=Document();doc.add_paragraph('保留原文');stream=io.BytesIO();doc.save(stream)
    source=legacy_manual_source(runtime,s['id'],stream.getvalue())
    # Exercise the analysis path independently of the existing manual-source guard.
    index=document_store.load(runtime,s['id'])
    index['documents'][0]['versions'][0]['source_type']='runtime'
    index['documents'][0]['versions'][0]['review_status']='accepted'
    (document_store.folder(runtime,s['id'])/'index.json').write_text(json.dumps(index))
    def runner(packet,emit,bridge,stop):
        bridge('route_request',{'domain':'disclosure','intent':'document','document_action':'revise',
            'target_document_id':source['document_id'],'reason':'局部修改'})
        item=draft(title='人工修改稿',document_id=source['document_id'],base_version=1,text='整篇重写')
        bridge('assess_document_readiness',{'documents':[readiness(item)],'request_quote':packet['prompt'],'decision':'assess'})
        with pytest.raises(HTTPException) as error:bridge('make_word',{'documents':[item]})
        assert error.value.status_code==409 and '局部' in error.value.detail
        emit({'type':'done'})
    runtime.runner=runner
    out=settled(c,request(c,s,'局部修改原Word'))
    assert out['run']['status']=='incomplete'
    assert listing(c,s)['items'][0]['version']==1
