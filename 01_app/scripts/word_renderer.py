"""External Agent Word tool. Not imported or executed by the Harness API."""
import io
import re
import copy
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from backend.document_content import parse, WordDeliveryError, required_header_fields
from backend.vendor.nero_word import word_primitives as wp


def render_document(packet, template_raw=None):
    """The runtime's document snapshot is independent of disclosure-stage approvals."""
    from backend.document_rendering import render as render_text
    from backend.document_files import patch
    import base64
    document=packet['document']
    if packet.get('source_base64'):
        return patch(base64.b64decode(packet['source_base64'],validate=True),document['edits'])[0]
    event={'company_name':packet.get('company_name',''),'title':document['title'],'layer':packet.get('board','chinext'),
           'plan':{'items':[]},'draft':{'text':document['text']}}
    if packet['layout'].get('template_authority')=='user_uploaded':
        return render_user_template(event,Document(io.BytesIO(template_raw)),parse(event),status_label='')[0]
    if document['kind']=='announcement':
        raw,manifest=render(event,packet['layout'],io.BytesIO(template_raw),status_label='')
        check=Document(io.BytesIO(raw))
        if [p.text for p in check.paragraphs]!=[p['text'] for p in manifest['paragraphs']] or \
                [[[c.text for c in r.cells] for r in t.rows] for t in check.tables]!=manifest['tables']:
            raise WordDeliveryError(422,'公告 Word 正文回读不一致，未交付文件')
        return raw
    return render_text(document['text'],document['title'],packet.get('company_name',''),packet['layout'],
                       packet['generated_at'],template_raw=template_raw,
                       notice=packet.get('notice','待审阅稿：依据当前资料整理，待补事项及内容、版式须复核；文件生成不代表事项获批或已发布。'))

def _font(run,p,size,bold=False):
    wp.apply_run_font(run,latin=p['fonts']['latin'],east_asia=p['fonts']['east_asia'],size_pt=size,bold=bold,color=RGBColor(0,0,0))

def _paragraph(doc,text,role,p,manifest):
    style={'title':'Title','h1':'Heading 1','h2':'Heading 2'}.get(role,'Normal')
    para=doc.add_paragraph(style=style);run=para.add_run(text)
    size=p['fonts'][{'title':'title_pt','h1':'heading_pt','h2':'heading_pt','meta':'meta_pt','status':'meta_pt'}.get(role,'body_pt')]
    bold=role in ('title','h1');_font(run,p,size,bold)
    fmt=para.paragraph_format;fmt.line_spacing=Pt(p['paragraph']['body_line_pt']);fmt.space_after=Pt(0);fmt.space_before=Pt(0)
    fmt.first_line_indent=Pt(p['paragraph']['first_line_indent_pt'] if role=='body' else 0)
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER if role=='title' else WD_ALIGN_PARAGRAPH.RIGHT if role=='signature' else WD_ALIGN_PARAGRAPH.LEFT
    if role=='title':fmt.space_after=Pt(12);wp.set_paragraph_keep(para,keep_next=True,keep_lines=True)
    if role in ('h1','h2'):
        fmt.space_before=Pt(p['paragraph']['heading_before_pt']);fmt.space_after=Pt(p['paragraph']['heading_after_pt']);wp.set_paragraph_keep(para,keep_next=True,keep_lines=True)
    if role in ('status','meta'):fmt.line_spacing=Pt(16)
    manifest.append({'text':text,'role':role,'size_pt':size,'bold':bold,'line_pt':16 if role in ('status','meta') else p['paragraph']['body_line_pt'],'indent_pt':p['paragraph']['first_line_indent_pt'] if role=='body' else 0,'alignment':int(para.alignment)})
    return para

def render(event,p,template,*,status_label='模拟验收稿'):
    """Thin builder over python-docx and shared primitives; no source Word is edited."""
    ir=parse(event);doc=Document(template)
    if p.get('template_authority')=='user_uploaded':return render_user_template(event,doc,ir,status_label=status_label)
    placeholders={'{{ title }}','{{ company_name }}','{{ status_label }}','{{ body_text }}','{{ text }}'}
    if any(x.text.strip() and x.text.strip() not in placeholders for x in doc.paragraphs):
        raise WordDeliveryError(409,'模板含非占位正文，请保护原稿并先核对模板用途')
    for el in list(doc.element.body):
        if el.tag.rsplit('}',1)[-1]!='sectPr':doc.element.body.remove(el)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        s=doc.styles[name]
        wp.apply_style_font(s,latin=p['fonts']['latin'],east_asia=p['fonts']['east_asia'],size_pt=p['fonts']['body_pt'],color=RGBColor(0,0,0))
        # Remove inherited theme overrides in these used styles; paragraph runs also carry explicit fonts.
        for node in s.element.iter():
            if node.tag.rsplit('}',1)[-1]=='rFonts':
                for key in list(node.attrib):
                    if 'theme' in key.lower():del node.attrib[key]
            if node.tag.rsplit('}',1)[-1]=='pBdr':node.getparent().remove(node)
    sec=doc.sections[0];spec=p['page'];sec.page_width=Mm(spec['width_mm']);sec.page_height=Mm(spec['height_mm'])
    for name in ('top','bottom','left','right'):setattr(sec,name+'_margin',Mm(spec[name+'_mm']))
    sec.header_distance=Mm(spec['header_mm']);sec.footer_distance=Mm(spec['footer_mm'])
    paragraphs=[];tables=[];blocks=list(ir['blocks']);leading=[]
    while blocks and blocks[0]['kind']=='meta':leading.append(blocks.pop(0))
    if status_label:_paragraph(doc,status_label, 'status',p,paragraphs)
    for block in leading:_paragraph(doc,block['text'],'meta',p,paragraphs)
    header_text='\n'.join(b['text'] for b in leading);missing=[]
    for key in required_header_fields(event):
        match=re.search(key+r'[：:](.*?)(?=(?:证券代码|证券简称|主办券商|公告编号)[：:]|$)',header_text,re.S)
        if not match or not match.group(1).strip() or any(x in match.group(1) for x in ('待补','待填写','待完善')):missing.append(key)
    if missing:_paragraph(doc,'【待补：'+ '、'.join(missing)+'】','meta',p,paragraphs)
    _paragraph(doc,event['company_name'],'title',p,paragraphs)
    title=ir['title'];title=title[len(event['company_name']):].strip() if title.startswith(event['company_name']) else title
    _paragraph(doc,title or '公告','title',p,paragraphs)
    for i,block in enumerate(blocks):
        if block['kind']!='table':
            role='h'+str(block['level']) if block['kind']=='heading' else 'meta' if block['kind']=='meta' else 'body'
            if role=='body' and i>=len(blocks)-3 and (block['text'].endswith('董事会') or re.fullmatch(r'\d{4}年\d{1,2}月\d{1,2}日',block['text'])):role='signature'
            _paragraph(doc,block['text'],role,p,paragraphs);continue
        values=[block['header']]+block['rows'];columns=len(block['header']);table=doc.add_table(rows=len(values),cols=columns);table.autofit=False
        width=int((sec.page_width-sec.left_margin-sec.right_margin)/columns)
        wp.set_table_borders(table,color=p['table']['border_color'],size=p['table']['border_size_eighth_pt'])
        for j,column in enumerate(table.columns):column.width=width
        for i,row in enumerate(table.rows):
            wp.set_row_cant_split(row,enabled=i==0);wp.set_repeat_table_header(row,enabled=i==0)
            for j,cell in enumerate(row.cells):
                cell.width=width;wp.set_cell_margins(cell,**{k:p['table']['cell_margin_dxa'] for k in ('top','bottom','start','end')})
                para=cell.paragraphs[0];para.paragraph_format.first_line_indent=Pt(0);para.paragraph_format.line_spacing=Pt(p['paragraph']['table_line_pt']);para.paragraph_format.space_after=Pt(0)
                para.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.RIGHT if re.fullmatch(r'[\d,.%+\-]+',values[i][j]) else WD_ALIGN_PARAGRAPH.LEFT
                _font(para.add_run(values[i][j]),p,p['fonts']['table_pt'],bold=i==0)
        tables.append(values)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in list(footer.runs):run._r.getparent().remove(run._r)
    run=wp.add_page_number(footer);_font(run,p,p['fonts']['footer_pt'])
    doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title=ir['title']
    output=io.BytesIO();doc.save(output)
    return output.getvalue(),{'paragraphs':paragraphs,'tables':tables,'missing_header_fields':missing}


def render_user_template(event,doc,ir,*,status_label='模拟验收稿'):
    """Fill anchors in a copy. Keep the user's sections, styles, furniture and other text."""
    paragraphs=list(doc.paragraphs)
    for table in doc.tables:
        paragraphs.extend(p for row in table.rows for cell in row.cells for p in cell.paragraphs)
    for section in doc.sections:
        if not section.header.is_linked_to_previous:paragraphs.extend(section.header.paragraphs)
        if not section.footer.is_linked_to_previous:paragraphs.extend(section.footer.paragraphs)
    body=next((p for p in paragraphs if re.fullmatch(r'\s*{{\s*body_text\s*}}\s*',p.text)),None)
    if body is None:raise WordDeliveryError(409,'手工模板需要独立的正文填充段落')
    blocks=copy.deepcopy(ir['blocks'])
    fixed=[p.text for p in doc.paragraphs if p.text.strip() and not re.search(r'{{.*?}}',p.text)]
    for text in fixed:
        index=next((i for i,b in enumerate(blocks) if b.get('text','').strip()==text.strip()),None)
        if index is None:raise WordDeliveryError(409,'模板固定正文尚未纳入已确认全文：'+text[:80])
        blocks.pop(index)
    replacements={'title':ir['title'],'company_name':event['company_name'],'status_label':status_label}
    for paragraph in paragraphs:
        if paragraph is body:continue
        for key,value in replacements.items():
            pattern=re.compile(r'{{\s*'+key+r'\s*}}')
            match=pattern.search(paragraph.text)
            if not match:continue
            # Replace across split runs while retaining unaffected run formatting.
            start,end=match.span();offset=0;inserted=False
            for run in paragraph.runs:
                a,b=offset,offset+len(run.text);offset=b
                if b<=start or a>=end:continue
                prefix=run.text[:max(0,start-a)];suffix=run.text[max(0,end-a):] if b>end else ''
                run.text=prefix+(value if not inserted else '')+suffix;inserted=True
    from docx.text.paragraph import Paragraph
    for block in blocks:
        if block['kind']=='table':
            values=[block['header']]+block['rows'];table=doc.add_table(rows=len(values),cols=len(values[0]))
            for r,row in enumerate(values):
                for c,value in enumerate(row):table.cell(r,c).text=value
            body._p.addprevious(table._tbl)
            wp.set_repeat_table_header(table.rows[0],enabled=True)
        else:
            element=copy.deepcopy(body._p)
            for child in list(element):
                if not child.tag.endswith('}pPr'):element.remove(child)
            paragraph=Paragraph(element,body._parent)
            run=paragraph.add_run(block['text'])
            if body.runs and body.runs[0]._r.rPr is not None:run._r.insert(0,copy.deepcopy(body.runs[0]._r.rPr))
            if block['kind']=='heading':paragraph.style='Heading '+str(min(block['level'],2))
            body._p.addprevious(element)
    body._p.getparent().remove(body._p)
    output=io.BytesIO();doc.save(output)
    return output.getvalue(),{'template_authority':'user_uploaded','paragraphs':[],'tables':[],'missing_header_fields':[]}
