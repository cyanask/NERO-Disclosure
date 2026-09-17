import re
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))
from backend import paths as workspace_paths
from backend.vendor.nero_word import word_primitives as wp

WORKSPACE = workspace_paths.repository_root(APP)
README_PATH = WORKSPACE / "README.md"
OUTPUT_DIR = workspace_paths.local_of(APP) / "output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DOCX = OUTPUT_DIR / "NERO_信息披露_AI系统说明.docx"

doc = Document()

# Page setup - A4 Portrait
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)
    
    # Header & Footer setup
    section.different_first_page_header_footer = False
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    wp.apply_run_font(f_p.add_run("— "), latin="Calibri", east_asia="SimSun", size_pt=9.0, color=RGBColor(0x88, 0x88, 0x88))
    wp.add_page_number(f_p, display_text="1")
    wp.apply_run_font(f_p.runs[-1], latin="Calibri", east_asia="SimSun", size_pt=9.0, color=RGBColor(0x88, 0x88, 0x88))
    wp.apply_run_font(f_p.add_run(" —"), latin="Calibri", east_asia="SimSun", size_pt=9.0, color=RGBColor(0x88, 0x88, 0x88))

# Color palette
COLOR_PRIMARY_NAVY = RGBColor(0x1F, 0x3A, 0x60)
COLOR_TEXT_MAIN = RGBColor(0x22, 0x22, 0x22)
COLOR_TEXT_MUTED = RGBColor(0x55, 0x55, 0x55)
HEX_HEADER_FILL = "1F3A60"
HEX_ROW_ALT = "F8FAFC"
HEX_BORDER = "D9D9D9"

# Helper for inline formatting
def add_formatted_runs(paragraph, text, base_size=10.5, is_bold=False, base_color=COLOR_TEXT_MAIN, latin_font="Calibri", east_asia_font="SimSun"):
    # Strip HTML comments if any
    text = re.sub(r'<!--.*?-->', '', text, flags=re.DOTALL)
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    if not text.strip():
        return

    # Parse tokens: bold, code, link, plain
    # Pattern to match **bold**, `code`, [link](url)
    pattern = re.compile(r'(\*{2}[^*]+?\*{2}|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))')
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            content = part[2:-2]
            run = paragraph.add_run(content)
            wp.apply_run_font(run, latin=latin_font, east_asia="SimHei", size_pt=base_size, bold=True, color=base_color)
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            content = part[1:-1]
            run = paragraph.add_run(content)
            wp.apply_run_font(run, latin="Consolas", east_asia="SimSun", size_pt=base_size * 0.92, bold=False, color=RGBColor(0x99, 0x33, 0x33))
        elif part.startswith('[') and '](' in part and part.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if m:
                link_text, link_url = m.group(1), m.group(2)
                run = paragraph.add_run(link_text)
                wp.apply_run_font(run, latin=latin_font, east_asia=east_asia_font, size_pt=base_size, bold=is_bold, color=COLOR_PRIMARY_NAVY)
            else:
                run = paragraph.add_run(part)
                wp.apply_run_font(run, latin=latin_font, east_asia=east_asia_font, size_pt=base_size, bold=is_bold, color=base_color)
        else:
            run = paragraph.add_run(part)
            wp.apply_run_font(run, latin=latin_font, east_asia=east_asia_font, size_pt=base_size, bold=is_bold, color=base_color)

# Read README
raw_lines = README_PATH.read_text(encoding="utf-8").splitlines()

# Filter out prose quality comment binding at the end
clean_lines = []
for line in raw_lines:
    if '<!-- prose-quality-binding:' in line:
        break
    clean_lines.append(line)

# Process line by line or block by block
i = 0
n = len(clean_lines)

while i < n:
    line = clean_lines[i].rstrip()
    
    # Blank line
    if not line:
        i += 1
        continue
    
    # Document Title (# )
    if line.startswith("# "):
        title_text = line[2:].strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(14)
        wp.set_paragraph_keep(p, keep_next=True)
        run = p.add_run(title_text)
        wp.apply_run_font(run, latin="Calibri", east_asia="SimHei", size_pt=22.0, bold=True, color=RGBColor(0x11, 0x11, 0x11))
        i += 1
        continue
    
    # Heading 1 (## )
    if line.startswith("## "):
        h1_text = line[3:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        wp.set_paragraph_keep(p, keep_next=True)
        run = p.add_run(h1_text)
        wp.apply_run_font(run, latin="Calibri", east_asia="SimHei", size_pt=14.0, bold=True, color=COLOR_PRIMARY_NAVY)
        i += 1
        continue
        
    # Heading 2 (### )
    if line.startswith("### "):
        h2_text = line[4:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        wp.set_paragraph_keep(p, keep_next=True)
        run = p.add_run(h2_text)
        wp.apply_run_font(run, latin="Calibri", east_asia="SimHei", size_pt=12.0, bold=True, color=RGBColor(0x22, 0x22, 0x22))
        i += 1
        continue
        
    # Blockquote (> )
    if line.startswith("> "):
        quote_text = line[2:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.5)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.25
        wp.set_paragraph_shading(p, "F4F6F8")
        add_formatted_runs(p, quote_text, base_size=10.0, base_color=RGBColor(0x33, 0x44, 0x55))
        i += 1
        continue
        
    # Markdown table (| ... |)
    if line.startswith("|") and line.endswith("|"):
        table_lines = []
        while i < n and clean_lines[i].strip().startswith("|") and clean_lines[i].strip().endswith("|"):
            table_lines.append(clean_lines[i].strip())
            i += 1
            
        # Parse table lines
        if len(table_lines) >= 2:
            headers = [c.strip() for c in table_lines[0].strip("|").split("|")]
            # Check if second line is separator
            is_sep = all(re.match(r'^:?-+:?$', c.strip()) for c in table_lines[1].strip("|").split("|"))
            start_row = 2 if is_sep else 1
            rows_data = []
            for t_line in table_lines[start_row:]:
                cols = [c.strip() for c in t_line.strip("|").split("|")]
                # Pad or trim to header length
                if len(cols) < len(headers):
                    cols += [""] * (len(headers) - len(cols))
                rows_data.append(cols[:len(headers)])
                
            num_cols = len(headers)
            table = doc.add_table(rows=1 + len(rows_data), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Compute column widths based on available width: 15.8 cm
            total_w = 15.8
            # Specific custom column distributions if matching known tables
            if num_cols == 3 and "电脑" in headers[0]:
                col_widths = [3.5, 4.8, 7.5]
            elif num_cols == 3 and "阶段" in headers[0]:
                col_widths = [3.8, 6.2, 5.8]
            elif num_cols == 3 and "方式" in headers[0]:
                col_widths = [3.8, 6.5, 5.5]
            elif num_cols == 2 and "入口" in headers[0]:
                col_widths = [3.8, 12.0]
            elif num_cols == 2 and "目录" in headers[0]:
                col_widths = [4.5, 11.3]
            else:
                col_widths = [total_w / num_cols] * num_cols
                
            # Set borders and header repeating
            wp.set_table_borders(table, color="D9D9D9", size=4, edges=("top", "bottom", "left", "right", "insideH", "insideV"))
            wp.set_repeat_table_header(table.rows[0], True)
            
            # Format header row
            hdr_row = table.rows[0]
            wp.set_row_cant_split(hdr_row, True)
            for c_idx, head_text in enumerate(headers):
                cell = hdr_row.cells[c_idx]
                cell.width = Cm(col_widths[c_idx])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                wp.set_cell_margins(cell, top=140, bottom=140, start=140, end=140)
                wp.set_cell_shading(cell, HEX_HEADER_FILL)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                if num_cols > 2 and c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(head_text)
                wp.apply_run_font(run, latin="Calibri", east_asia="SimHei", size_pt=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                
            # Format data rows
            for r_idx, row_vals in enumerate(rows_data):
                row = table.rows[1 + r_idx]
                wp.set_row_cant_split(row, True)
                fill_color = HEX_ROW_ALT if (r_idx % 2 == 1) else "FFFFFF"
                for c_idx, val_text in enumerate(row_vals):
                    cell = row.cells[c_idx]
                    cell.width = Cm(col_widths[c_idx])
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    wp.set_cell_margins(cell, top=100, bottom=100, start=140, end=140)
                    if fill_color != "FFFFFF":
                        wp.set_cell_shading(cell, fill_color)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.2
                    if num_cols > 2 and c_idx == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    add_formatted_runs(p, val_text, base_size=9.5, base_color=COLOR_TEXT_MAIN)
                    
            # Add small spacing after table
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(6)
            continue
            
    # Bullet list item (- ...)
    if line.startswith("- "):
        bullet_text = line[2:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        run_bullet = p.add_run("•  ")
        wp.apply_run_font(run_bullet, latin="Calibri", east_asia="SimSun", size_pt=10.5, bold=True, color=COLOR_PRIMARY_NAVY)
        add_formatted_runs(p, bullet_text, base_size=10.5, base_color=COLOR_TEXT_MAIN)
        i += 1
        continue
        
    # Numbered list item (1. ..., 2. ...)
    num_match = re.match(r'^(d+).s+(.*)', line)
    if num_match:
        num_str = num_match.group(1)
        item_text = num_match.group(2).strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        run_num = p.add_run(f"{num_str}.  ")
        wp.apply_run_font(run_num, latin="Calibri", east_asia="SimHei", size_pt=10.5, bold=True, color=COLOR_PRIMARY_NAVY)
        add_formatted_runs(p, item_text, base_size=10.5, base_color=COLOR_TEXT_MAIN)
        
        # Check if there are indented paragraphs following this numbered item
        i += 1
        while i < n and (clean_lines[i].startswith("   ") or clean_lines[i].startswith("	")):
            sub_line = clean_lines[i].strip()
            if sub_line:
                sub_p = doc.add_paragraph()
                sub_p.paragraph_format.left_indent = Cm(1.4)
                sub_p.paragraph_format.space_before = Pt(2)
                sub_p.paragraph_format.space_after = Pt(3)
                sub_p.paragraph_format.line_spacing = 1.25
                add_formatted_runs(sub_p, sub_line, base_size=10.0, base_color=COLOR_TEXT_MAIN)
            i += 1
        continue
        
    # Normal body paragraph
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.28
    add_formatted_runs(p, line, base_size=10.5, base_color=COLOR_TEXT_MAIN)
    i += 1

doc.save(OUTPUT_DOCX)
print(f"SUCCESS: Saved {OUTPUT_DOCX}")
