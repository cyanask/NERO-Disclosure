"""Receive and verify external Word output. This module never creates a Word file."""
import base64
import hashlib
import io
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from uuid import uuid4
from docx import Document
from fastapi import HTTPException
from .document_content import parse, required_header_fields


def normalized(value):return re.sub(r'\s+','',value)


def check(raw,event,layout,template_raw=None):
    errors=[]
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries=archive.infolist()
            if len(entries)>2000 or sum(i.file_size for i in entries)>30000000:raise ValueError('Word 解压后过大')
            if any('embeddings/' in i.filename.lower() or 'vbaproject' in i.filename.lower() for i in entries):raise ValueError('Word 不得携带嵌入对象或宏')
            for entry in entries:
                if entry.filename.endswith('.rels'):
                    relations=ET.fromstring(archive.read(entry))
                    if any(node.get('TargetMode','').casefold()=='external' for node in relations):raise ValueError('Word 含外部关系，请移除后重新登记')
        doc=Document(io.BytesIO(raw));ir=parse(event)
    except Exception as exc:
        raise HTTPException(422,'Word 文件结构无效或不受支持：'+str(exc)[:200])
    if layout.get('template_authority')=='user_uploaded':
        if not template_raw:raise HTTPException(409,'手工模板校验缺少当前原文件')
        return check_user_template(raw,doc,event,ir,template_raw)
    blocks=list(ir['blocks']);leading=[]
    while blocks and blocks[0]['kind']=='meta':leading.append(blocks.pop(0)['text'])
    title=ir['title'];title=title[len(event['company_name']):].strip() if title.startswith(event['company_name']) else title
    expected=['模拟验收稿',*leading,event['company_name'],title or '公告']
    for block in blocks:
        if block['kind']=='table':
            expected.extend(block['header'])
            for row in block['rows']:expected.extend(row)
        else:expected.append(block['text'])
    actual=[]
    for node in doc.element.body:
        if node.tag.endswith('}p'):actual.append(''.join(node.xpath('.//w:t/text()')))
        elif node.tag.endswith('}tbl'):
            actual.extend(''.join(cell.xpath('.//w:t/text()')) for cell in node.xpath('./w:tr/w:tc'))
    if [normalized(x) for x in actual if x.strip()] != [normalized(x) for x in expected if x.strip()]:
        errors.append('Word 正文或表格与当前 canonical draft.text 不一致')
    header='\n'.join(leading)
    for key in required_header_fields(event):
        match=re.search(key+r'[：:](.*?)(?=(?:证券代码|证券简称|主办券商|公告编号)[：:]|$)',header,re.S)
        if not match or not match.group(1).strip() or any(x in match.group(1) for x in ('待补','待填写','待完善')):errors.append('公告基础标识缺失：'+key)
    for sec in doc.sections:
        spec=layout['page']
        values={'width_mm':sec.page_width.mm,'height_mm':sec.page_height.mm,**{n+'_mm':getattr(sec,n+'_margin').mm for n in ('top','bottom','left','right')}}
        if any(abs(value-spec[key])>.1 for key,value in values.items()):errors.append('纸张或页边距不符合当前版式规则')
    paragraphs=list(doc.paragraphs)+[p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
    allowed_sizes={float(v) for k,v in layout['fonts'].items() if k.endswith('_pt')}
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():continue
            fonts=run._r.xpath('./w:rPr/w:rFonts');east=fonts[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia') if fonts else None
            if east!=layout['fonts']['east_asia'] or run.font.name!=layout['fonts']['latin'] or run.font.size is None or run.font.size.pt not in allowed_sizes:
                errors.append('字体或字号不符合当前版式规则');break
    return {'status':'FAIL' if errors else 'PASS','errors':list(dict.fromkeys(errors)),
        'checks':['DOCX容器及外部关系','正文与表格逐项一致性','基础标识','纸张与页边距','字体与字号'],
        'visual_acceptance':'not_verified','legal_acceptance':'not_verified'}


def template_bytes(seeds,event):
    """Read the same scoped, hash-bound template for registration and recheck."""
    template=next((t for t in seeds.for_event(event).templates() if t['id']==event['draft']['template_id']),None)
    if not template:raise HTTPException(409,'当前手工模板不存在')
    root=(seeds.root/'templates').resolve();path=root/template['file']
    if path.is_symlink() or not path.resolve().is_relative_to(root):raise HTTPException(409,'手工模板路径无效')
    try:raw=path.read_bytes()
    except OSError:raise HTTPException(409,'手工模板文件不可读') from None
    if hashlib.sha256(raw).hexdigest()!=event['draft'].get('template_sha256'):raise HTTPException(409,'手工模板与已确认版本不一致')
    return raw


def path(directory,artifact):
    base=(Path(directory)/'artifacts').resolve();target=base/artifact['file']
    if target.is_symlink() or not target.resolve().is_relative_to(base) or not target.is_file():raise HTTPException(409,'交付物缺失或路径无效')
    if hashlib.sha256(target.read_bytes()).hexdigest()!=artifact['sha256']:raise HTTPException(409,'交付物哈希不符')
    return target


def check_user_template(raw,doc,event,ir,template_raw):
    """Independent content and format comparison against draft text and the uploaded original."""
    baseline=Document(io.BytesIO(template_raw));errors=[]
    def flat(document):
        values=[]
        for node in document.element.body:
            if node.tag.endswith('}p'):values.append(''.join(node.xpath('.//w:t/text()')))
            elif node.tag.endswith('}tbl'):values.extend(''.join(cell.xpath('.//w:t/text()')) for cell in node.xpath('./w:tr/w:tc'))
        return [normalized(v) for v in values if v.strip()]
    actual=flat(doc);expected=[]
    for block in ir['blocks']:
        if block['kind']=='table':
            expected.extend(normalized(x) for x in block['header'])
            for row in block['rows']:expected.extend(normalized(x) for x in row)
        else:expected.append(normalized(block['text']))
    values={'title':ir['title'],'company_name':event['company_name'],'status_label':'模拟验收稿'}
    for para in baseline.paragraphs:
        if '{{' not in para.text or 'body_text' in para.text:continue
        filled=para.text
        for key,value in values.items():filled=re.sub(r'{{\s*'+key+r'\s*}}',lambda _:value,filled)
        try:actual.remove(normalized(filled))
        except ValueError:errors.append('模板标题或公司标识填充不一致')
    if actual!=expected:errors.append('Word 正文、固定内容或表格与已确认全文不一致')
    def xml(raw):return ET.tostring(ET.fromstring(raw))
    def properties(paragraph):
        return (xml(paragraph._p.pPr.xml.encode()) if paragraph._p.pPr is not None else None,
                [xml(r._r.rPr.xml.encode()) if r._r.rPr is not None else None for r in paragraph.runs])
    for paragraph in baseline.paragraphs:
        if 'body_text' in paragraph.text:continue
        filled=paragraph.text
        for key,value in values.items():filled=re.sub(r'{{\s*'+key+r'\s*}}',lambda _:value,filled)
        matches=[p for p in doc.paragraphs if p.text==filled]
        if matches and not any(properties(p)==properties(paragraph) for p in matches):errors.append('模板保留段落的字体或段落格式被改写')
    with zipfile.ZipFile(io.BytesIO(raw)) as out,zipfile.ZipFile(io.BytesIO(template_raw)) as original:
        for name in original.namelist():
            if name.startswith(('word/styles','word/theme/','word/media/','word/fontTable','word/numbering')):
                if name not in out.namelist() or (xml(out.read(name))!=xml(original.read(name)) if name.endswith('.xml') else out.read(name)!=original.read(name)):
                    errors.append('手工模板样式或素材被改写：'+name)
    if len(doc.sections)!=len(baseline.sections) or any(xml(a._sectPr.xml.encode())!=xml(b._sectPr.xml.encode()) for a,b in zip(doc.sections,baseline.sections)):
        errors.append('手工模板纸张、页边距或分节设置变化')
    for actual_section,original_section in zip(doc.sections,baseline.sections):
        for a,b in ((actual_section.header,original_section.header),(actual_section.footer,original_section.footer)):
            expected_text='\n'.join(p.text for p in b.paragraphs)
            for key,value in values.items():expected_text=re.sub(r'{{\s*'+key+r'\s*}}',lambda _:value,expected_text)
            if '\n'.join(p.text for p in a.paragraphs)!=expected_text:errors.append('手工页眉或页脚内容被改写')
            if [properties(p) for p in a.paragraphs]!=[properties(p) for p in b.paragraphs]:errors.append('手工页眉或页脚格式被改写')
    anchor=next((p for p in baseline.paragraphs if re.fullmatch(r'\s*{{\s*body_text\s*}}\s*',p.text)),None)
    if anchor:
        expected_rpr=xml(anchor.runs[0]._r.rPr.xml.encode()) if anchor.runs and anchor.runs[0]._r.rPr is not None else None
        baseline_text={p.text for p in baseline.paragraphs}
        for paragraph in doc.paragraphs:
            if normalized(paragraph.text) not in expected or paragraph.text in baseline_text:continue
            for run in paragraph.runs:
                actual_rpr=xml(run._r.rPr.xml.encode()) if run._r.rPr is not None else None
                if run.text.strip() and actual_rpr!=expected_rpr:errors.append('正文填充字体格式偏离手工模板')
    return {'status':'FAIL' if errors else 'PASS','errors':list(dict.fromkeys(errors)),
        'checks':['DOCX容器及外部关系','已确认全文与表格顺序','上传模板样式与素材','分节与页边距','手工页眉页脚','正文填充格式'],
        'visual_acceptance':'not_verified','legal_acceptance':'not_verified','template_authority':'user_uploaded'}
