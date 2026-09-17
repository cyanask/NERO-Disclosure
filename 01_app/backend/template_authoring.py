"""Create reviewable template candidates from existing registered layout assets."""
import copy,hashlib,json,re
from pathlib import Path
from fastapi import HTTPException
from .domain import Seeds
from . import paths as workspace_paths
from .library_admin import atomic,locked


def replacement(root, board, import_id, profile_id, company='', expected=None, apply=False):
    """Keep the exact uploaded bytes as the current version, never rebuild the user's file."""
    from . import file_ingestion, announcement_history
    from docx import Document
    root=Path(root);source=file_ingestion.get(root,import_id,board,'profiles',company)
    path=workspace_paths.resolve(root,source['original_path'])
    if path.suffix!='.docx':raise HTTPException(422,'可编辑模板须为 DOCX 文件')
    doc=Document(path)
    all_paragraphs=list(doc.paragraphs)+[p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
    if sum(bool(re.fullmatch(r'\s*{{\s*body_text\s*}}\s*',p.text)) for p in all_paragraphs)!=1:
        raise HTTPException(422,'模板需要且只能有一个独立的 {{ body_text }} 正文填充段落；请保留原稿并核对该位置')
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    missing=[key for key in ('title','body_text') if '{{ '+key+' }}' not in text and '{{'+key+'}}' not in text]
    if missing:raise HTTPException(422,'模板缺少填充位置：'+ '、'.join(missing)+'；请保留手工稿并补回这些位置')
    seed=Seeds(root).for_board(board);common=json.loads((seed.template_dir/'manifest.json').read_text())
    base=next((r for r in common if r.get('profile_id')==profile_id),None)
    if not base:raise HTTPException(422,'请先选择要替换的模板')
    target_dir=root/'templates/companies'/board/announcement_history.company_key(company) if company else seed.template_dir
    manifest=target_dir/'manifest.json'
    with locked(root,board):
        old=manifest.read_bytes() if manifest.exists() else b'[]\n';rows=json.loads(old)
        current=next((r for r in rows if r['id']==base['id']),base)
        fingerprint=hashlib.sha256(old).hexdigest()
        preview={'template_id':base['id'],'title':base['name'],'company':company,'fingerprint':fingerprint,
            'old_sha256':hashlib.sha256((root/'templates'/current['file']).read_bytes()).hexdigest(),
            'new_sha256':source['sha256'],'message':'以手工文件作为当前模板，已有任务需重新核对模板版本；不改写已交付文件。'}
        if not apply:return preview
        if expected!=fingerprint:raise HTTPException(409,'模板版本已变化，请重新预览')
        target_dir.mkdir(parents=True,exist_ok=True);history=target_dir/'history';history.mkdir(exist_ok=True)
        atomic(history/(fingerprint+'.json'),old)
        output=target_dir/(source['sha256']+'.docx');atomic(output,path.read_bytes())
        entry={**current,'file':str(output.relative_to(root/'templates')),'version':source['sha256'][:12],
            'sha256':source['sha256'],'authority':'user_uploaded','updated_at':announcement_history.now(),'company_scope':company or None}
        rows=[r for r in rows if r['id']!=entry['id']]+[entry]
        atomic(manifest,(json.dumps(rows,ensure_ascii=False,indent=2)+'\n').encode())
        return {**preview,'status':'applied','file':entry['file'],'authority':'user_uploaded'}


def add_uploaded(root,board,import_id,base_profile_id,title,company=''):
    from . import file_ingestion,library_admin
    source=file_ingestion.get(root,import_id,board,'profiles',company)
    if not title.strip():raise HTTPException(422,'请填写模板名称')
    # Validate the uploaded file before registering any content profile.
    replacement(root,board,import_id,base_profile_id,company)
    state=library_admin.state(root,'profiles',board)
    original=next((r for r in state['items'] if r['id']==base_profile_id),None)
    if not original:raise HTTPException(404,'参考文种不存在')
    profile_id='profile-upload-'+source['sha256'][:20]
    if not any(r['id']==profile_id for r in state['items']):
        profile={**copy.deepcopy(original),'id':profile_id,'title':title.strip(),'company_scope':company or None,'case_evidence':[],
            'scope_review_status':'format_and_case_evidence_pending','review_status':'user_uploaded_pending_review','verification_status':'candidate_unreviewed'}
        library_admin.update(root,'profiles',[profile],state['fingerprint'],board)
    # Existing profile-to-template fallback supplies the registered base style.
    seed=Seeds(root,board,{'layer':board,'stock_code':company})
    base=next(r for r in seed.templates() if r.get('profile_id')==profile_id)
    base={**base,'company_scope':company or None}
    manifest=seed.template_dir/'manifest.json'
    with locked(root,board):
        old=manifest.read_bytes();rows=json.loads(old)
        if not any(r.get('profile_id')==profile_id for r in rows):
            archive=seed.template_dir/'history';archive.mkdir(exist_ok=True)
            atomic(archive/(hashlib.sha256(old).hexdigest()+'.json'),old)
            atomic(manifest,(json.dumps(rows+[base],ensure_ascii=False,indent=2)+'\n').encode())
    preview=replacement(root,board,import_id,profile_id,company)
    return replacement(root,board,import_id,profile_id,company,preview['fingerprint'],True)


def prepare_template(root,rid,board,profile):
    from docx import Document
    root=Path(root);seeds=Seeds(root).for_board(board)
    manifest=json.loads((seeds.template_dir/'manifest.json').read_text())
    base=next((x for x in manifest if x.get('layout_profile_id')==profile['layout_profile_id']),None)
    if not base:raise HTTPException(409,'当前版式没有已登记基础模板；请先核对可复用版式')
    original=(root/'templates'/base['file']).resolve()
    if not original.is_relative_to((root/'templates').resolve()) or not original.is_file():raise HTTPException(409,'基础模板文件不可用')
    output=workspace_paths.work(root)/'template-candidates'/rid/(profile['id']+'.docx');output.parent.mkdir(parents=True,exist_ok=True)
    doc=Document(original)
    for child in list(doc.element.body):
        if not child.tag.endswith('}sectPr'):doc.element.body.remove(child)
    doc.add_paragraph('{{ title }}','Title');doc.add_paragraph('{{ company_name }}')
    doc.add_paragraph('{{ status_label }}');doc.add_paragraph('{{ body_text }}')
    doc.core_properties.title=profile['title'];doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.save(output)
    value={**copy.deepcopy(base),'id':'tpl-'+profile['id'],'name':profile['title'],'profile_id':profile['id'],'kind':profile['kind'],
           'file':'generated/'+board+'/'+profile['id']+'.docx','sections':profile['sections'],'version':'1.0.0',
           'scope_review_status':'format_and_case_evidence_pending','review_status':'pending_human_review','library_board':board,
           'source_note':('依据登记法源及参考案例制作的模板候选；' if profile.get('case_evidence') else '依据登记法源制作，未使用案例参考；')+'内容与版式待人工验收。'}
    return {'candidate_path':workspace_paths.store_path(root,output),'sha256':hashlib.sha256(output.read_bytes()).hexdigest(),
            'manifest_entry':value,'base_template_id':base['id'],'human_review':'pending'}


def register_template(root,board,proposal):
    root=Path(root);seeds=Seeds(root).for_board(board);entry=proposal['manifest_entry']
    path=workspace_paths.resolve(root,proposal['candidate_path']).resolve()
    if not path.is_relative_to((workspace_paths.work(root)/'template-candidates').resolve()):raise HTTPException(403,'模板候选路径越界')
    raw=path.read_bytes()
    if hashlib.sha256(raw).hexdigest()!=proposal['sha256']:raise HTTPException(409,'模板候选文件已变化')
    with locked(root,board):
        manifest=seeds.template_dir/'manifest.json';old=manifest.read_bytes();rows=json.loads(old)
        existing=next((r for r in rows if r['id']==entry['id']),None)
        if existing:
            existing_file=root/'templates'/existing['file']
            if existing!=entry or not existing_file.is_file() or hashlib.sha256(existing_file.read_bytes()).hexdigest()!=proposal['sha256']:raise HTTPException(409,'模板编号已经存在且内容不同')
            return {'id':entry['id'],'file':entry['file'],'sha256':proposal['sha256'],'review_status':'pending_human_review'}
        target=root/'templates'/entry['file'];target.parent.mkdir(parents=True,exist_ok=True)
        if target.exists():raise HTTPException(409,'模板目标文件已经存在')
        atomic(target,raw)
        history=seeds.template_dir/'history';history.mkdir(exist_ok=True)
        atomic(history/('manifest-'+hashlib.sha256(old).hexdigest()+'.json'),old)
        atomic(manifest,(json.dumps(rows+[entry],ensure_ascii=False,indent=2)+'\n').encode())
    return {'id':entry['id'],'file':entry['file'],'sha256':proposal['sha256'],'review_status':'pending_human_review'}
