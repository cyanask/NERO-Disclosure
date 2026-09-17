"""Shared Word rendering primitives. Persistence and authorization belong to runtime."""
import hashlib
import io
import json
import re
import time
from pathlib import Path
from uuid import uuid4
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor
from fastapi import HTTPException
from .document_content import WordDeliveryError, cells
from .vendor.nero_word import word_primitives as wp
from . import paths as workspace_paths

MAX_TEXT = 200000
MAX_BLOCKS = 600
KIND = 'consult_work_draft'
NOTICE = '咨询工作稿／待复核：来自信披咨询讨论，未经正式披露流程核验，不构成公告或法律意见。'


def layout(root, board):
    """Reuse the board's retained typography; a missing profile falls back to A4 defaults."""
    path = workspace_paths.templates(root) / 'boards' / board / 'layout_profiles.json'
    try:
        profiles = json.loads(path.read_text('utf-8'))
    except (OSError, ValueError):
        profiles = []
    for row in profiles:
        if row.get('page') and row.get('fonts'):
            return row
    return {'page': {'width_mm': 210, 'height_mm': 297, 'top_mm': 25.4, 'bottom_mm': 25.4
                     , 'left_mm': 31.75, 'right_mm': 31.75, 'header_mm': 12.7, 'footer_mm': 12.7},
            'fonts': {'east_asia': '宋体', 'latin': 'Times New Roman', 'body_pt': 12, 'title_pt': 14
                      , 'heading_pt': 12, 'table_pt': 10.5, 'meta_pt': 10.5, 'footer_pt': 9},
            'paragraph': {'body_line_pt': 24, 'first_line_indent_pt': 24, 'heading_before_pt': 12
                          , 'heading_after_pt': 6, 'table_line_pt': 14},
            'table': {'border_color': '000000', 'border_size_eighth_pt': 4, 'cell_margin_dxa': 80}}


def blocks_from_text(text):
    """Derived block model. The displayed reply stays the canonical content."""
    lines = text.replace('\r\n', '\n').split('\n')
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith('```'):
            body = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append({'kind': 'paragraph', 'text': '\n'.join(body) if any(body) else '（空代码块）'})
            continue
        if line.startswith('|') and i < len(lines) and lines[i].strip().startswith('|') \
                and all(re.fullmatch(r':?-{3,}:?', c) for c in cells(lines[i])):
            header = cells(line)
            raw_lines = [line, lines[i]]
            separator = cells(lines[i])
            i += 1
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = cells(lines[i])
                raw_lines.append(lines[i])
                i += 1
                rows.append(row)
            if rows and len(separator) == len(header) and all(len(row) == len(header) for row in rows):
                blocks.append({'kind': 'table', 'header': header, 'rows': rows})
                continue
            # A malformed table is still user-confirmed content. Retain every
            # line verbatim rather than guessing column ownership or dropping rows.
            blocks.append({'kind': 'paragraph', 'text': '\n'.join(raw_lines)})
            continue
        heading = re.match(r'^(#{1,6})\s+(.*)$', line)
        if heading:
            blocks.append({'kind': 'heading', 'level': min(len(heading.group(1)), 3),
                           'text': clean(heading.group(2))})
            continue
        if re.match(r'^\d{1,3}[.、)]\s+', line):
            blocks.append({'kind': 'paragraph', 'text': clean(line)})
            continue
        if re.match(r'^[-*+]\s+', line):
            blocks.append({'kind': 'bullet', 'text': clean(re.sub(r'^[-*+]\s+', '', line))})
            continue
        if line.startswith('>'):
            blocks.append({'kind': 'quote', 'text': clean(line.lstrip('> ').strip())})
            continue
        if line.startswith('---') and set(line) <= set('-—'):
            continue
        blocks.append({'kind': 'paragraph', 'text': clean(line)})
    if len(blocks) > MAX_BLOCKS:
        raise HTTPException(422, '本轮内容过长，请分次导出')
    return blocks


def clean(value):
    value = re.sub(r'\*\*(.+?)\*\*', r'\1', value)
    value = re.sub(r'(?<!\w)\*(?!\s)(.+?)(?<!\s)\*(?!\w)', r'\1', value)
    value = re.sub(r'`([^`]*)`', r'\1', value)
    # Keep destinations in the file. Complex Markdown links remain verbatim.
    value = re.sub(r'\[([^\]]+)\]\(([^()\n]+)\)', r'\1（\2）', value)
    return value.strip()


def title_of(text, fallback):
    """Prefer an explicit Markdown heading, then the session title.

    The reply's first sentence is often a lead-in such as “可以。” and makes a
    poor document title and filename.
    """
    for line in text.replace('\r\n', '\n').split('\n'):
        value = line.strip()
        if value.startswith('#'):
            heading = clean(re.sub(r'^#{1,6}\s+', '', value))
            if heading:
                return heading[:120]
    session = (fallback or '').strip()
    if session and session != '新会话':
        return session[:120]
    for line in text.replace('\r\n', '\n').split('\n'):
        value = clean(line.strip())
        if value:
            return value[:120]
    return '咨询工作稿'


def render(text, title, company, board_layout, generated_at, *, template_raw=None, notice=NOTICE):
    spec = board_layout
    doc = Document(io.BytesIO(template_raw)) if template_raw else Document()
    if template_raw:
        placeholders = {'{{ title }}', '{{ company_name }}', '{{ status_label }}', '{{ body_text }}', '{{ text }}'}
        if doc.tables or any(p.text.strip() and p.text.strip() not in placeholders for p in doc.paragraphs):
            raise WordDeliveryError(409, '模板含固定正文，请选择适用模板或保留原稿进行锚定修改')
        for node in list(doc.element.body):
            if not node.tag.endswith('}sectPr'):
                doc.element.body.remove(node)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        try:
            wp.apply_style_font(doc.styles[name], latin=spec['fonts']['latin'],
                                east_asia=spec['fonts']['east_asia'],
                                size_pt=spec['fonts']['body_pt'], color=RGBColor(0, 0, 0))
        except KeyError:
            continue
    section = doc.sections[0]
    page = spec['page']
    section.page_width = Mm(page['width_mm'])
    section.page_height = Mm(page['height_mm'])
    for name in ('top', 'bottom', 'left', 'right'):
        setattr(section, name + '_margin', Mm(page[name + '_mm']))
    section.header_distance = Mm(page['header_mm'])
    section.footer_distance = Mm(page['footer_mm'])
    body = spec['paragraph']
    fonts = spec['fonts']

    def paragraph(value, role='body', size=None, bold=False, indent=None, align=None, keep=False):
        para = doc.add_paragraph(style='Title' if role == 'title' else None)
        run = para.add_run(value)
        wp.apply_run_font(run, latin=fonts['latin'], east_asia=fonts['east_asia'],
                          size_pt=size or fonts['body_pt'], bold=bold, color=RGBColor(0, 0, 0))
        fmt = para.paragraph_format
        fmt.line_spacing = Pt(body['table_line_pt'] if role == 'meta' else body['body_line_pt'])
        fmt.space_before = Pt(body['heading_before_pt'] if role.startswith('h') else 0)
        fmt.space_after = Pt(body['heading_after_pt'] if role.startswith('h') else 0)
        fmt.first_line_indent = Pt(body['first_line_indent_pt'] if indent is None and role == 'body' else (indent or 0))
        para.alignment = align or (WD_ALIGN_PARAGRAPH.CENTER if role == 'title' else WD_ALIGN_PARAGRAPH.LEFT)
        if keep:
            wp.set_paragraph_keep(para, keep_next=True, keep_lines=True)
        return para

    paragraph('文稿 · 待审阅' if notice != NOTICE else '咨询工作稿 · 待复核', 'meta', size=fonts['meta_pt'], bold=True
              , align=WD_ALIGN_PARAGRAPH.CENTER)
    if company:
        paragraph(company, 'meta', size=fonts['meta_pt'])
    paragraph('生成时间：' + generated_at, 'meta', size=fonts['meta_pt'])
    paragraph(notice, 'meta', size=fonts['meta_pt'])
    blocks = blocks_from_text(text)
    leading_title = bool(blocks and blocks[0]['kind'] == 'heading' and blocks[0]['text'] == title)
    if not leading_title:
        paragraph(title, 'title', size=fonts['title_pt'], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
    prefix = [p.text for p in doc.paragraphs]
    for block_index, block in enumerate(blocks):
        if block['kind'] == 'table':
            values = [block['header']] + block['rows']
            columns = max(len(row) for row in values)
            table = doc.add_table(rows=len(values), cols=columns)
            table.autofit = False
            width = int((section.page_width - section.left_margin - section.right_margin) / columns)
            wp.set_table_borders(table, color=spec['table']['border_color'],
                                 size=spec['table']['border_size_eighth_pt'])
            for column in table.columns:
                column.width = width
            for index, row in enumerate(table.rows):
                wp.set_row_cant_split(row, enabled=index == 0)
                wp.set_repeat_table_header(row, enabled=index == 0)
                for position, cell in enumerate(row.cells):
                    cell.width = width
                    wp.set_cell_margins(cell, **{k: spec['table']['cell_margin_dxa']
                                                 for k in ('top', 'bottom', 'start', 'end')})
                    para = cell.paragraphs[0]
                    para.paragraph_format.first_line_indent = Pt(0)
                    para.paragraph_format.line_spacing = Pt(body['table_line_pt'])
                    para.paragraph_format.space_after = Pt(0)
                    value = values[index][position] if position < len(values[index]) else ''
                    run = para.add_run(value)
                    wp.apply_run_font(run, latin=fonts['latin'], east_asia=fonts['east_asia'],
                                      size_pt=fonts['table_pt'], bold=index == 0, color=RGBColor(0, 0, 0))
            continue
        if block['kind'] == 'heading':
            if leading_title and block_index == 0:
                paragraph(block['text'], 'title', size=fonts['title_pt'], bold=True,
                          align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
            else:
                paragraph(block['text'], 'h' + str(block['level']),
                          size=fonts['heading_pt'], bold=True, indent=0, keep=True)
        elif block['kind'] == 'bullet':
            paragraph('· ' + block['text'], 'body', indent=body['first_line_indent_pt'])
        elif block['kind'] == 'quote':
            paragraph(block['text'], 'meta', size=fonts['meta_pt'], indent=body['first_line_indent_pt'])
        else:
            for part in str(block['text']).split('\n'):
                paragraph(part, 'body')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in list(footer.runs):
        run._r.getparent().remove(run._r)
    run = wp.add_page_number(footer)
    wp.apply_run_font(run, latin=fonts['latin'], east_asia=fonts['east_asia'],
                      size_pt=fonts['footer_pt'], color=RGBColor(0, 0, 0))
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = '咨询工作稿 · ' + title
    output = io.BytesIO()
    doc.save(output)
    raw = output.getvalue()
    verify_rendered_content(raw, prefix, blocks)
    return raw


def verify_rendered_content(raw, prefix, blocks):
    """Read back OOXML in document order before allowing a file to be delivered."""
    expected = [('paragraph', value) for value in prefix]
    for block in blocks:
        if block['kind'] == 'table':
            expected.append(('table', [block['header']] + block['rows']))
        else:
            value = ('· ' if block['kind'] == 'bullet' else '') + block['text']
            expected.extend(('paragraph', part) for part in value.split('\n'))
    doc = Document(io.BytesIO(raw))
    actual = []
    for node in doc.element.body:
        if node.tag.endswith('}p'):
            actual.append(('paragraph', Paragraph(node, doc).text))
        elif node.tag.endswith('}tbl'):
            table = Table(node, doc)
            actual.append(('table', [[cell.text for cell in row.cells] for row in table.rows]))
    if actual != expected:
        raise WordDeliveryError(422, 'Word 内容回读与已确认正文不一致，未交付文件，请检查转换结果')
